"""Deterministic, read-only S1 audit for the 2024 C problem attachments.

The four XLSX files are opened with openpyxl in read-only mode.  All derived
artifacts are written below results/raw/s1; no workbook is ever saved.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import struct
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import numpy as np
from openpyxl import load_workbook


SCRIPT_VERSION = "1.0.0"
WAVEFORM_POINTS = 1024
EXPECTED_HASHES = {
    "problem/数据驱动下磁性元件的磁芯损耗建模.docx": "3268B0B7A1A3CCC5E72CC1490F7389063A7DF35E35D307228E952504054E1EC8",
    "src/附件一（训练集）.xlsx": "C16AEE712BC0AD480B5285F17FA7FA57A5E22A1901D4A6ADB0A1EB83ECFB7FFC",
    "src/附件二（测试集）.xlsx": "B46A4BF69331A41D4CF6190C97906914619C76335BEF43EF37BA148CD99FC66C",
    "src/附件三（测试集）.xlsx": "8C403B8D7717E42567EC7C6DEC204FEFF87023B83D45B87EC0C983C53A2D929A",
    "src/附件四（Excel表）.xlsx": "D8FDFFDF63839F7B40D5DD87BE6923AF04AD1D52180072A47C5B4DF70832C1C8",
}
TRAIN_SHEETS = {
    "材料1": (3401, 1028),
    "材料2": (3001, 1028),
    "材料3": (3201, 1028),
    "材料4": (2801, 1028),
}
ALLOWED_TEMPERATURES = {25.0, 50.0, 70.0, 90.0}
ALLOWED_MATERIALS = set(TRAIN_SHEETS)
ALLOWED_WAVEFORMS = {"正弦波", "三角波", "梯形波"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def finite_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def encode_value(value: Any) -> bytes:
    if value is None:
        return b"N"
    if isinstance(value, bool):
        return b"B1" if value else b"B0"
    if isinstance(value, (int, float, np.integer, np.floating)):
        number = float(value)
        if math.isnan(number):
            return b"Q"
        if math.isinf(number):
            return b"P" if number > 0 else b"M"
        return b"F" + struct.pack("<d", number)
    encoded = str(value).encode("utf-8")
    return b"S" + len(encoded).to_bytes(4, "little") + encoded


def fingerprint(values: Iterable[Any]) -> str:
    digest = hashlib.sha256()
    for value in values:
        encoded = encode_value(value)
        digest.update(len(encoded).to_bytes(4, "little"))
        digest.update(encoded)
    return digest.hexdigest()


def valid_waveform_header(values: tuple[Any, ...]) -> bool:
    if len(values) != WAVEFORM_POINTS or not str(values[0]).startswith("0"):
        return False
    return all(finite_number(value) == float(index) for index, value in enumerate(values[1:], start=1))


@dataclass
class PositionStats:
    count: np.ndarray
    missing: np.ndarray
    nonfinite: np.ndarray
    total: np.ndarray
    total_sq: np.ndarray
    minimum: np.ndarray
    maximum: np.ndarray

    @classmethod
    def create(cls) -> "PositionStats":
        return cls(
            count=np.zeros(WAVEFORM_POINTS, dtype=np.int64),
            missing=np.zeros(WAVEFORM_POINTS, dtype=np.int64),
            nonfinite=np.zeros(WAVEFORM_POINTS, dtype=np.int64),
            total=np.zeros(WAVEFORM_POINTS, dtype=np.float64),
            total_sq=np.zeros(WAVEFORM_POINTS, dtype=np.float64),
            minimum=np.full(WAVEFORM_POINTS, np.inf, dtype=np.float64),
            maximum=np.full(WAVEFORM_POINTS, -np.inf, dtype=np.float64),
        )

    def update(self, values: tuple[Any, ...]) -> tuple[np.ndarray, np.ndarray]:
        if len(values) != WAVEFORM_POINTS:
            raise ValueError(f"Expected {WAVEFORM_POINTS} waveform points, got {len(values)}")
        missing_mask = np.fromiter((value is None for value in values), dtype=bool, count=WAVEFORM_POINTS)
        numeric = np.empty(WAVEFORM_POINTS, dtype=np.float64)
        for index, value in enumerate(values):
            try:
                numeric[index] = float(value) if value is not None else np.nan
            except (TypeError, ValueError):
                numeric[index] = np.nan
        finite_mask = np.isfinite(numeric)
        self.count += finite_mask
        self.missing += missing_mask
        self.nonfinite += (~finite_mask & ~missing_mask)
        safe = np.where(finite_mask, numeric, 0.0)
        self.total += safe
        self.total_sq += safe * safe
        self.minimum = np.minimum(self.minimum, np.where(finite_mask, numeric, np.inf))
        self.maximum = np.maximum(self.maximum, np.where(finite_mask, numeric, -np.inf))
        return numeric, finite_mask

    def rows(self, dataset: str, scope: str) -> Iterable[dict[str, Any]]:
        for index in range(WAVEFORM_POINTS):
            count = int(self.count[index])
            if count:
                mean = self.total[index] / count
                variance = max(self.total_sq[index] / count - mean * mean, 0.0)
                minimum = float(self.minimum[index])
                maximum = float(self.maximum[index])
                std = math.sqrt(variance)
            else:
                mean = minimum = maximum = std = None
            yield {
                "dataset": dataset,
                "scope": scope,
                "position": index,
                "finite_count": count,
                "missing_count": int(self.missing[index]),
                "nonfinite_count": int(self.nonfinite[index]),
                "min": minimum,
                "max": maximum,
                "mean": mean,
                "std_population": std,
            }


class Audit:
    def __init__(self, project_root: Path, output_dir: Path) -> None:
        self.project_root = project_root
        self.output_dir = output_dir
        self.dataset_rows: list[dict[str, Any]] = []
        self.quality_rows: list[dict[str, Any]] = []
        self.categories: Counter[tuple[str, str, str, str]] = Counter()
        self.numeric: defaultdict[tuple[str, str, str], list[float]] = defaultdict(list)
        self.numeric_missing: Counter[tuple[str, str, str]] = Counter()
        self.numeric_nonfinite: Counter[tuple[str, str, str]] = Counter()
        self.position_stats: dict[tuple[str, str], PositionStats] = {}
        self.wave_hashes: defaultdict[str, Counter[str]] = defaultdict(Counter)
        self.row_hashes: defaultdict[str, Counter[str]] = defaultdict(Counter)
        self.input_hashes_before: dict[str, str] = {}
        self.input_hashes_after: dict[str, str] = {}

    def check(
        self,
        check_id: str,
        dataset: str,
        scope: str,
        severity: str,
        passed: bool,
        value: Any,
        expected: Any,
        notes: str,
    ) -> None:
        self.quality_rows.append(
            {
                "check_id": check_id,
                "dataset": dataset,
                "scope": scope,
                "severity": severity,
                "status": "PASS" if passed else ("WARN" if severity in {"INFO", "RISK"} else "FAIL"),
                "value": value,
                "expected": expected,
                "notes": notes,
            }
        )

    def add_numeric(self, dataset: str, scope: str, field: str, value: Any) -> None:
        key = (dataset, scope, field)
        if value is None:
            self.numeric_missing[key] += 1
            return
        try:
            number = float(value)
        except (TypeError, ValueError):
            self.numeric_nonfinite[key] += 1
            return
        if not math.isfinite(number):
            self.numeric_nonfinite[key] += 1
            return
        self.numeric[key].append(number)

    def add_category(self, dataset: str, scope: str, field: str, value: Any) -> None:
        rendered = "<MISSING>" if value is None else str(value)
        self.categories[(dataset, scope, field, rendered)] += 1

    def update_waveform(
        self,
        dataset: str,
        scope: str,
        waveform: tuple[Any, ...],
    ) -> str:
        position_stats = self.position_stats.setdefault((dataset, scope), PositionStats.create())
        numeric, finite_mask = position_stats.update(waveform)
        finite = numeric[finite_mask]
        if finite.size:
            metrics = {
                "b_peak_abs_T": float(np.max(np.abs(finite))),
                "b_peak_to_peak_T": float(np.max(finite) - np.min(finite)),
                "b_mean_T": float(np.mean(finite)),
                "b_rms_T": float(np.sqrt(np.mean(finite * finite))),
            }
            for field, value in metrics.items():
                self.add_numeric(dataset, scope, field, value)
        else:
            for field in ("b_peak_abs_T", "b_peak_to_peak_T", "b_mean_T", "b_rms_T"):
                self.add_numeric(dataset, scope, field, None)
        wave_hash = fingerprint(waveform)
        self.wave_hashes[f"{dataset}/{scope}"][wave_hash] += 1
        return wave_hash

    def update_training_all(self, scope: str, metadata: tuple[Any, ...], waveform: tuple[Any, ...]) -> None:
        temperature, frequency, loss, wave_type = metadata
        for field, value in (
            ("temperature_C", temperature),
            ("frequency_Hz", frequency),
            ("core_loss_W_per_m3", loss),
        ):
            self.add_numeric("train", "all", field, value)
        self.add_category("train", "all", "material", scope)
        self.add_category("train", "all", "temperature_C", temperature)
        self.add_category("train", "all", "waveform", wave_type)
        self.add_category("train", "all", "material_x_temperature_x_waveform", f"{scope}|{temperature}|{wave_type}")
        wave_hash = self.update_waveform("train", "all", waveform)
        self.row_hashes["train"][fingerprint((scope, *metadata, wave_hash))] += 1

    def audit_train(self, path: Path) -> None:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            self.check(
                "TRAIN_SHEETS",
                "train",
                "workbook",
                "CRITICAL",
                workbook.sheetnames == list(TRAIN_SHEETS),
                "|".join(workbook.sheetnames),
                "|".join(TRAIN_SHEETS),
                "工作表名称与顺序。",
            )
            for sheet_name, expected_shape in TRAIN_SHEETS.items():
                if sheet_name not in workbook.sheetnames:
                    raise ValueError(f"Missing training sheet: {sheet_name}")
                sheet = workbook[sheet_name]
                actual_shape = (sheet.max_row, sheet.max_column)
                self.check(
                    "TRAIN_SHAPE",
                    "train",
                    sheet_name,
                    "CRITICAL",
                    actual_shape == expected_shape,
                    f"{actual_shape[0]}x{actual_shape[1]}",
                    f"{expected_shape[0]}x{expected_shape[1]}",
                    "含表头；实际文件为 4 个元字段 + 1024 点。",
                )
                rows = sheet.iter_rows(values_only=True)
                header = tuple(next(rows))
                expected_prefix = ("温度，oC", "频率，Hz", "磁芯损耗，w/m3", "励磁波形")
                self.check(
                    "TRAIN_HEADER_PREFIX",
                    "train",
                    sheet_name,
                    "CRITICAL",
                    header[:4] == expected_prefix,
                    "|".join(map(str, header[:4])),
                    "|".join(expected_prefix),
                    "元数据字段顺序。",
                )
                self.check(
                    "TRAIN_WAVEFORM_WIDTH",
                    "train",
                    sheet_name,
                    "CRITICAL",
                    len(header[4:]) == WAVEFORM_POINTS,
                    len(header[4:]),
                    WAVEFORM_POINTS,
                    "等间距单周期磁通密度采样点数。",
                )
                header_valid = valid_waveform_header(header[4:])
                self.check("TRAIN_WAVEFORM_HEADER", "train", sheet_name, "CRITICAL", header_valid, "0..1023" if header_valid else "INVALID", "0..1023", "采样位置表头连续且顺序正确。")
                row_count = 0
                invalid_wave_width = 0
                local_rows: Counter[str] = Counter()
                for row in rows:
                    row_count += 1
                    metadata = tuple(row[:4])
                    waveform = tuple(row[4:])
                    if len(waveform) != WAVEFORM_POINTS:
                        invalid_wave_width += 1
                        continue
                    temperature, frequency, loss, wave_type = metadata
                    for field, value in (
                        ("temperature_C", temperature),
                        ("frequency_Hz", frequency),
                        ("core_loss_W_per_m3", loss),
                    ):
                        self.add_numeric("train", sheet_name, field, value)
                    self.add_category("train", sheet_name, "waveform", wave_type)
                    self.add_category("train", sheet_name, "temperature_C", temperature)
                    wave_hash = self.update_waveform("train", sheet_name, waveform)
                    local_rows[fingerprint((*metadata, wave_hash))] += 1
                    self.update_training_all(sheet_name, metadata, waveform)
                self.row_hashes[f"train/{sheet_name}"] = local_rows
                self.check("TRAIN_ROW_COUNT", "train", sheet_name, "CRITICAL", row_count == expected_shape[0] - 1, row_count, expected_shape[0] - 1, "不含表头。")
                self.check("TRAIN_WAVEFORM_ROW_WIDTH", "train", sheet_name, "CRITICAL", invalid_wave_width == 0, invalid_wave_width, 0, "波形长度异常的记录数。")
                self.dataset_rows.append(
                    {
                        "file": path.name,
                        "dataset": "train",
                        "scope": sheet_name,
                        "role": "labeled_training",
                        "rows": row_count,
                        "columns": sheet.max_column,
                        "waveform_points": WAVEFORM_POINTS,
                        "label_fields": "core_loss_W_per_m3|waveform",
                        "id_field": "",
                    }
                )
            total_rows = sum(row["rows"] for row in self.dataset_rows if row["dataset"] == "train")
            self.dataset_rows.append(
                {
                    "file": path.name,
                    "dataset": "train",
                    "scope": "all",
                    "role": "labeled_training",
                    "rows": total_rows,
                    "columns": 1029,
                    "waveform_points": WAVEFORM_POINTS,
                    "label_fields": "core_loss_W_per_m3|waveform",
                    "id_field": "",
                }
            )
        finally:
            workbook.close()

    def audit_test(
        self,
        path: Path,
        dataset: str,
        expected_shape: tuple[int, int],
        metadata_width: int,
        expected_prefix: tuple[str, ...],
    ) -> None:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            self.check("TEST_SHEETS", dataset, "workbook", "CRITICAL", workbook.sheetnames == ["测试集"], "|".join(workbook.sheetnames), "测试集", "测试工作表。")
            if "测试集" not in workbook.sheetnames:
                raise ValueError(f"Missing 测试集 sheet in {path.name}")
            sheet = workbook["测试集"]
            actual_shape = (sheet.max_row, sheet.max_column)
            self.check("TEST_SHAPE", dataset, "测试集", "CRITICAL", actual_shape == expected_shape, f"{actual_shape[0]}x{actual_shape[1]}", f"{expected_shape[0]}x{expected_shape[1]}", "含表头。")
            rows = sheet.iter_rows(values_only=True)
            header = tuple(next(rows))
            self.check("TEST_HEADER_PREFIX", dataset, "测试集", "CRITICAL", header[:metadata_width] == expected_prefix, "|".join(map(str, header[:metadata_width])), "|".join(expected_prefix), "元数据字段顺序。")
            self.check("TEST_WAVEFORM_WIDTH", dataset, "测试集", "CRITICAL", len(header[metadata_width:]) == WAVEFORM_POINTS, len(header[metadata_width:]), WAVEFORM_POINTS, "等间距单周期磁通密度采样点数。")
            header_valid = valid_waveform_header(header[metadata_width:])
            self.check("TEST_WAVEFORM_HEADER", dataset, "测试集", "CRITICAL", header_valid, "0..1023" if header_valid else "INVALID", "0..1023", "采样位置表头连续且顺序正确。")
            ids: list[int] = []
            row_count = 0
            invalid_wave_width = 0
            for row in rows:
                row_count += 1
                metadata = tuple(row[:metadata_width])
                waveform = tuple(row[metadata_width:])
                if len(waveform) != WAVEFORM_POINTS:
                    invalid_wave_width += 1
                    continue
                sample_id, temperature, frequency, material = metadata[:4]
                if finite_number(sample_id) is not None:
                    ids.append(int(float(sample_id)))
                for field, value in (("sample_id", sample_id), ("temperature_C", temperature), ("frequency_Hz", frequency)):
                    self.add_numeric(dataset, "测试集", field, value)
                self.add_category(dataset, "测试集", "material", material)
                self.add_category(dataset, "测试集", "temperature_C", temperature)
                if metadata_width == 5:
                    self.add_category(dataset, "测试集", "waveform", metadata[4])
                wave_hash = self.update_waveform(dataset, "测试集", waveform)
                self.row_hashes[dataset][fingerprint((*metadata[1:], wave_hash))] += 1
            expected_rows = expected_shape[0] - 1
            self.check("TEST_ROW_COUNT", dataset, "测试集", "CRITICAL", row_count == expected_rows, row_count, expected_rows, "不含表头。")
            self.check("TEST_ID_UNIQUE", dataset, "测试集", "CRITICAL", len(ids) == len(set(ids)) == expected_rows, len(ids) - len(set(ids)), 0, "重复样本序号数。")
            self.check("TEST_ID_RANGE", dataset, "测试集", "CRITICAL", sorted(ids) == list(range(1, expected_rows + 1)), f"{min(ids) if ids else None}..{max(ids) if ids else None}", f"1..{expected_rows}", "样本序号完整性。")
            self.check("TEST_WAVEFORM_ROW_WIDTH", dataset, "测试集", "CRITICAL", invalid_wave_width == 0, invalid_wave_width, 0, "波形长度异常的记录数。")
            self.dataset_rows.append(
                {
                    "file": path.name,
                    "dataset": dataset,
                    "scope": "测试集",
                    "role": "classification_test" if dataset == "test2" else "loss_prediction_test",
                    "rows": row_count,
                    "columns": sheet.max_column,
                    "waveform_points": WAVEFORM_POINTS,
                    "label_fields": "" if dataset == "test2" else "waveform_input_only",
                    "id_field": "sample_id",
                }
            )
        finally:
            workbook.close()

    def audit_template(self, path: Path) -> None:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            expected_sheets = {"Sheet1", "Sheet2", "Sheet3"}
            self.check("TEMPLATE_SHEETS", "template4", "workbook", "CRITICAL", set(workbook.sheetnames) == expected_sheets, "|".join(workbook.sheetnames), "Sheet1|Sheet2|Sheet3", "附件四工作表集合。")
            sheet = workbook["Sheet1"]
            rows = list(sheet.iter_rows(values_only=True))
            ids = [int(float(row[0])) for row in rows if row and finite_number(row[0]) is not None]
            data_rows = rows[1:]
            col2_nonblank = sum(1 for row in data_rows if len(row) > 1 and row[1] is not None)
            col3_nonblank = sum(1 for row in data_rows if len(row) > 2 and row[2] is not None)
            self.check("TEMPLATE_SHAPE", "template4", "Sheet1", "CRITICAL", (sheet.max_row, sheet.max_column) == (401, 3), f"{sheet.max_row}x{sheet.max_column}", "401x3", "400 个样本的统一结果表。")
            expected_header = ("序号", "附件二（80个样品）励磁波形分类结果", "附件三（400个样品）磁芯损耗预测结果")
            self.check("TEMPLATE_HEADER", "template4", "Sheet1", "CRITICAL", tuple(rows[0]) == expected_header, "|".join(map(str, rows[0])), "|".join(expected_header), "附件四输出列名称与顺序。")
            self.check("TEMPLATE_ID_RANGE", "template4", "Sheet1", "CRITICAL", ids == list(range(1, 401)), f"{min(ids) if ids else None}..{max(ids) if ids else None}", "1..400", "首列样本序号。")
            self.check("TEMPLATE_OUTPUT_BLANK", "template4", "Sheet1", "INFO", col2_nonblank == 0 and col3_nonblank == 0, f"col2={col2_nonblank};col3={col3_nonblank}", "col2=0;col3=0", "S1 只检查容量，不填预测。")
            self.dataset_rows.append(
                {
                    "file": path.name,
                    "dataset": "template4",
                    "scope": "Sheet1",
                    "role": "submission_template",
                    "rows": 400,
                    "columns": sheet.max_column,
                    "waveform_points": 0,
                    "label_fields": "q1_classification_slot|q4_loss_prediction_slot",
                    "id_field": "sample_id",
                }
            )
        finally:
            workbook.close()

    @staticmethod
    def duplicate_rows(counter: Counter[str]) -> tuple[int, int]:
        groups = sum(1 for count in counter.values() if count > 1)
        excess_rows = sum(count - 1 for count in counter.values() if count > 1)
        return groups, excess_rows

    def add_semantic_checks(self) -> None:
        for dataset, scopes in (("train", [*TRAIN_SHEETS, "all"]), ("test2", ["测试集"]), ("test3", ["测试集"])):
            for scope in scopes:
                temp_values = self.numeric[(dataset, scope, "temperature_C")]
                frequency_values = self.numeric[(dataset, scope, "frequency_Hz")]
                self.check("TEMPERATURE_DOMAIN", dataset, scope, "MAJOR", bool(temp_values) and set(temp_values).issubset(ALLOWED_TEMPERATURES), "|".join(map(lambda x: f"{x:g}", sorted(set(temp_values)))), "25|50|70|90", "题目声明的离散温度。")
                outside_frequency = sum(value < 50000 or value > 500000 for value in frequency_values)
                self.check("FREQUENCY_DOMAIN", dataset, scope, "RISK", bool(frequency_values) and outside_frequency == 0, f"range={min(frequency_values):g}..{max(frequency_values):g};outside={outside_frequency}" if frequency_values else "EMPTY", "50000..500000;outside=0", "实测频率对题目文字范围存在轻微越界时保留原值，并在后续切分/建模中作为数据范围风险处理。")
                if dataset == "train":
                    wave_values = {value for (ds, sc, field, value), count in self.categories.items() if ds == dataset and sc == scope and field == "waveform" and count}
                    self.check("WAVEFORM_DOMAIN", dataset, scope, "MAJOR", wave_values.issubset(ALLOWED_WAVEFORMS) and bool(wave_values), "|".join(sorted(wave_values)), "正弦波|三角波|梯形波", "训练标签取值。")
                    loss_values = self.numeric[(dataset, scope, "core_loss_W_per_m3")]
                    self.check("LOSS_POSITIVE", dataset, scope, "MAJOR", bool(loss_values) and min(loss_values) > 0, min(loss_values) if loss_values else "EMPTY", ">0", "磁芯损耗密度应为正。")
                if dataset in {"test2", "test3"}:
                    material_values = {value for (ds, sc, field, value), count in self.categories.items() if ds == dataset and sc == scope and field == "material" and count}
                    self.check("MATERIAL_DOMAIN", dataset, scope, "MAJOR", material_values.issubset(ALLOWED_MATERIALS) and bool(material_values), "|".join(sorted(material_values)), "材料1|材料2|材料3|材料4", "测试集材料取值。")

        for name, counter in sorted(self.row_hashes.items()):
            groups, excess = self.duplicate_rows(counter)
            self.check("EXACT_CONTENT_DUPLICATES", name.split("/")[0], name, "INFO", True, f"groups={groups};excess_rows={excess}", "diagnostic", "测试集内容重复检查不含样本序号；重复不自动判为错误。")
        for name, counter in sorted(self.wave_hashes.items()):
            groups, excess = self.duplicate_rows(counter)
            dataset, scope = name.split("/", 1)
            self.check("EXACT_WAVEFORM_DUPLICATES", dataset, scope, "INFO", True, f"groups={groups};excess_rows={excess}", "diagnostic", "基于 1024 个 IEEE-754 数值的确定性指纹；重复不自动判为错误。")

        train_hashes = set(self.wave_hashes["train/all"])
        for dataset in ("test2", "test3"):
            test_counter = self.wave_hashes[f"{dataset}/测试集"]
            overlap_hashes = train_hashes & set(test_counter)
            overlap_rows = sum(test_counter[key] for key in overlap_hashes)
            self.check("TRAIN_TEST_WAVEFORM_OVERLAP", dataset, "all", "MAJOR", overlap_rows == 0, f"hashes={len(overlap_hashes)};test_rows={overlap_rows}", 0, "与训练集 1024 点波形完全相同会带来泄漏/近重复风险。")
        overlap_23 = set(self.wave_hashes["test2/测试集"]) & set(self.wave_hashes["test3/测试集"])
        self.check("TEST2_TEST3_WAVEFORM_OVERLAP", "test2_vs_test3", "all", "INFO", True, len(overlap_23), "diagnostic", "两个官方测试集之间的完全相同波形指纹数。")

        for key in sorted(set(self.numeric) | set(self.numeric_missing) | set(self.numeric_nonfinite)):
            dataset, scope, field = key
            missing = self.numeric_missing[key]
            nonfinite = self.numeric_nonfinite[key]
            self.check("NUMERIC_COMPLETENESS", dataset, scope, "MAJOR", missing == 0 and nonfinite == 0, f"{field}:missing={missing};nonfinite={nonfinite}", 0, "数值字段完整性。")
        for (dataset, scope), stats in sorted(self.position_stats.items()):
            missing = int(stats.missing.sum())
            nonfinite = int(stats.nonfinite.sum())
            finite_lengths_ok = bool(np.all(stats.count == stats.count[0]))
            self.check("WAVEFORM_COMPLETENESS", dataset, scope, "MAJOR", missing == 0 and nonfinite == 0 and finite_lengths_ok, f"missing={missing};nonfinite={nonfinite};uniform_finite_count={finite_lengths_ok}", 0, "全部 1024 点应为有限数值。")
            p2p = self.numeric[(dataset, scope, "b_peak_to_peak_T")]
            flat_count = sum(value <= 0 for value in p2p)
            self.check("FLAT_WAVEFORMS", dataset, scope, "MAJOR", flat_count == 0, flat_count, 0, "峰峰值为零的波形数。")

    def numeric_rows(self) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        keys = sorted(set(self.numeric) | set(self.numeric_missing) | set(self.numeric_nonfinite))
        for dataset, scope, field in keys:
            values = np.asarray(self.numeric[(dataset, scope, field)], dtype=np.float64)
            if values.size:
                q1, median, q3 = np.quantile(values, [0.25, 0.5, 0.75])
                iqr = q3 - q1
                lower, upper = q1 - 1.5 * iqr, q3 + 1.5 * iqr
                outlier_count = int(np.sum((values < lower) | (values > upper)))
                row = {
                    "dataset": dataset,
                    "scope": scope,
                    "field": field,
                    "finite_count": int(values.size),
                    "missing_count": self.numeric_missing[(dataset, scope, field)],
                    "nonfinite_count": self.numeric_nonfinite[(dataset, scope, field)],
                    "min": float(np.min(values)),
                    "q1": float(q1),
                    "median": float(median),
                    "q3": float(q3),
                    "max": float(np.max(values)),
                    "mean": float(np.mean(values)),
                    "std_population": float(np.std(values)),
                    "iqr_outlier_count": outlier_count,
                }
            else:
                row = {
                    "dataset": dataset,
                    "scope": scope,
                    "field": field,
                    "finite_count": 0,
                    "missing_count": self.numeric_missing[(dataset, scope, field)],
                    "nonfinite_count": self.numeric_nonfinite[(dataset, scope, field)],
                    "min": None,
                    "q1": None,
                    "median": None,
                    "q3": None,
                    "max": None,
                    "mean": None,
                    "std_population": None,
                    "iqr_outlier_count": 0,
                }
            rows.append(row)
        return rows

    @staticmethod
    def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)

    def run(self) -> None:
        for relative, expected_hash in EXPECTED_HASHES.items():
            path = self.project_root / relative
            if not path.is_file():
                raise FileNotFoundError(f"Required input missing: {relative}")
            actual_hash = sha256_file(path)
            self.input_hashes_before[relative] = actual_hash
            if actual_hash != expected_hash:
                raise ValueError(f"SHA-256 mismatch before audit: {relative}: {actual_hash} != {expected_hash}")

        self.audit_train(self.project_root / "src/附件一（训练集）.xlsx")
        self.audit_test(
            self.project_root / "src/附件二（测试集）.xlsx",
            "test2",
            (81, 1028),
            4,
            ("序号", "温度，oC", "频率，Hz", "磁芯材料"),
        )
        self.audit_test(
            self.project_root / "src/附件三（测试集）.xlsx",
            "test3",
            (401, 1029),
            5,
            ("序号", "温度，oC", "频率，Hz", "磁芯材料", "励磁波形"),
        )
        self.audit_template(self.project_root / "src/附件四（Excel表）.xlsx")
        self.add_semantic_checks()

        for relative, expected_hash in EXPECTED_HASHES.items():
            actual_hash = sha256_file(self.project_root / relative)
            self.input_hashes_after[relative] = actual_hash
            if actual_hash != expected_hash or actual_hash != self.input_hashes_before[relative]:
                raise RuntimeError(f"Input changed during audit: {relative}")

        self.output_dir.mkdir(parents=True, exist_ok=True)
        numeric_rows = self.numeric_rows()
        category_rows = [
            {"dataset": ds, "scope": scope, "field": field, "value": value, "count": count}
            for (ds, scope, field, value), count in sorted(self.categories.items())
        ]
        position_rows = [
            row
            for (dataset, scope), stats in sorted(self.position_stats.items())
            for row in stats.rows(dataset, scope)
        ]
        self.write_csv(
            self.output_dir / "dataset_summary.csv",
            self.dataset_rows,
            ["file", "dataset", "scope", "role", "rows", "columns", "waveform_points", "label_fields", "id_field"],
        )
        self.write_csv(
            self.output_dir / "categorical_counts.csv",
            category_rows,
            ["dataset", "scope", "field", "value", "count"],
        )
        self.write_csv(
            self.output_dir / "numeric_ranges.csv",
            numeric_rows,
            ["dataset", "scope", "field", "finite_count", "missing_count", "nonfinite_count", "min", "q1", "median", "q3", "max", "mean", "std_population", "iqr_outlier_count"],
        )
        self.write_csv(
            self.output_dir / "quality_checks.csv",
            self.quality_rows,
            ["check_id", "dataset", "scope", "severity", "status", "value", "expected", "notes"],
        )
        self.write_csv(
            self.output_dir / "waveform_position_profile.csv",
            position_rows,
            ["dataset", "scope", "position", "finite_count", "missing_count", "nonfinite_count", "min", "max", "mean", "std_population"],
        )

        failed = [row for row in self.quality_rows if row["status"] == "FAIL"]
        warnings = [row for row in self.quality_rows if row["status"] == "WARN"]
        critical_major_failed = [row for row in failed if row["severity"] in {"CRITICAL", "MAJOR"}]
        profile = {
            "script_version": SCRIPT_VERSION,
            "execution": {
                "environment": "math_modeling",
                "mode": "read_only_openpyxl_data_only",
                "deterministic": True,
                "model_training_performed": False,
                "official_test_sets_used_for_tuning": False,
            },
            "input_hashes_before": self.input_hashes_before,
            "input_hashes_after": self.input_hashes_after,
            "dataset_summary": self.dataset_rows,
            "quality_summary": {
                "checks": len(self.quality_rows),
                "failed": len(failed),
                "warnings": len(warnings),
                "critical_or_major_failed": len(critical_major_failed),
                "failed_ids": [f"{row['check_id']}:{row['dataset']}:{row['scope']}" for row in failed],
                "warning_ids": [f"{row['check_id']}:{row['dataset']}:{row['scope']}" for row in warnings],
            },
            "leakage_policy": {
                "attachment_2": "prediction_only_after_model_freeze; never for feature selection, tuning, or model selection",
                "attachment_3": "prediction_only_after_model_freeze; never for feature selection, tuning, or model selection",
                "exact_overlap_check": "SHA-256 over typed serialization of all 1024 waveform values",
            },
            "notes": [
                "IQR outlier counts are diagnostics, not automatic deletion rules.",
                "Test data were structurally audited only; no labels were inferred and no model was fit.",
                "Attachment 4 was opened read-only and was not filled or saved.",
            ],
        }
        with (self.output_dir / "data_profile.json").open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(profile, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
        if critical_major_failed:
            raise SystemExit(f"Audit produced {len(critical_major_failed)} CRITICAL/MAJOR failed checks; inspect quality_checks.csv")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    default_project = Path(__file__).resolve().parents[1]
    parser.add_argument("--project-root", type=Path, default=default_project)
    parser.add_argument("--output-dir", type=Path, default=None)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    project_root = args.project_root.resolve()
    output_dir = args.output_dir
    if output_dir is None:
        output_dir = project_root / "results/raw/s1"
    elif not output_dir.is_absolute():
        output_dir = project_root / output_dir
    Audit(project_root, output_dir.resolve()).run()
    print(f"S1 audit PASS: {output_dir.resolve()}")


if __name__ == "__main__":
    main()
