"""Generate the one-time S5 predictions from the G4-approved frozen models."""

from __future__ import annotations

import argparse
import json
import math
import subprocess
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from build_features import WAVEFORM_POINTS, waveform_features, waveform_sha256
from s3_common import (
    CLASS_ENCODING,
    DEFAULT_CONFIG,
    MATERIAL_LABELS,
    PROJECT_ROOT,
    SHAPE_FEATURES,
    TEMPERATURES,
    TZ,
    WAVEFORM_LABELS,
    ExperimentRun,
    load_frozen_config,
    set_global_seed,
    sha256_file,
    verify_contract_registry,
    verify_raw_inputs,
    write_csv,
    write_json,
)
from s4_common import predict_log_model


EXPERIMENT_ID = "EXP-S5-PRED-001"
DESCRIPTOR = PROJECT_ROOT / "experiments" / "main" / f"{EXPERIMENT_ID}.json"
RAW_S5_ROOT = PROJECT_ROOT / "results" / "raw" / "s5"
RUN_ROOT = RAW_S5_ROOT / "runs"
TEST2_FILE = PROJECT_ROOT / "src" / "附件二（测试集）.xlsx"
TEST3_FILE = PROJECT_ROOT / "src" / "附件三（测试集）.xlsx"
TEMPLATE_FILE = PROJECT_ROOT / "src" / "附件四（Excel表）.xlsx"
Q1_MODEL_FILE = PROJECT_ROOT / "results" / "raw" / "s3" / "q1" / "logistic_full_model.joblib"
Q1_METRICS_FILE = PROJECT_ROOT / "results" / "raw" / "s3" / "q1" / "baseline_metrics.json"
Q1_RUN_MANIFEST = PROJECT_ROOT / "results" / "raw" / "baseline" / "EXP-Q1-BASE-001" / "run_manifest.json"
Q4_MODEL_FILE = PROJECT_ROOT / "results" / "raw" / "main" / "q4" / "final_full_model.joblib"
Q4_WINNER_FILE = PROJECT_ROOT / "results" / "raw" / "main" / "q4" / "final_winner.json"
Q4_METRICS_FILE = PROJECT_ROOT / "results" / "raw" / "main" / "q4" / "ablation_metrics.json"
Q4_RUN_MANIFEST = PROJECT_ROOT / "results" / "raw" / "main" / "runs" / "EXP-Q4-ABL-001" / "run_manifest.json"
G4_REVIEW_FILE = PROJECT_ROOT / "reviews" / "gate_4_review.md"
G4_REVIEWED_COMMIT = "2309b1e361701be9818544fa18740cd8f5694f2e"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    return parser.parse_args()


def review_commit(path: Path) -> str:
    relative = path.relative_to(PROJECT_ROOT.parents[1]).as_posix()
    result = subprocess.run(
        ["git", "-C", str(PROJECT_ROOT.parents[1]), "log", "-1", "--format=%H", "--", relative],
        check=True,
        capture_output=True,
        text=True,
    )
    commit = result.stdout.strip()
    if len(commit) != 40:
        raise ValueError(f"Cannot resolve review commit for {relative}")
    return commit


def assert_g4_pass() -> dict[str, str]:
    text = G4_REVIEW_FILE.read_text(encoding="utf-8")
    required = (
        "- Verdict: **PASS**",
        f"- Reviewed Commit: `{G4_REVIEWED_COMMIT}`",
        "S4 -> S5 = AUTHORIZED",
    )
    missing = [token for token in required if token not in text]
    if missing:
        raise ValueError(f"G4 authorization is incomplete: {missing}")
    return {
        "review_file": "reviews/gate_4_review.md",
        "review_file_sha256": sha256_file(G4_REVIEW_FILE),
        "review_commit": review_commit(G4_REVIEW_FILE),
        "reviewed_commit": G4_REVIEWED_COMMIT,
        "verdict": "PASS",
    }


def load_test_features(path: Path, dataset: str, plateau_fraction: float) -> pd.DataFrame:
    if dataset == "test2":
        expected_rows, metadata_width = 80, 4
        expected_prefix = ("序号", "温度，oC", "频率，Hz", "磁芯材料")
    elif dataset == "test3":
        expected_rows, metadata_width = 400, 5
        expected_prefix = ("序号", "温度，oC", "频率，Hz", "磁芯材料", "励磁波形")
    else:
        raise ValueError(f"Unknown dataset: {dataset}")

    source_file_hash = sha256_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    try:
        if workbook.sheetnames != ["测试集"]:
            raise ValueError(f"Unexpected sheets in {path.name}: {workbook.sheetnames}")
        sheet = workbook["测试集"]
        if (sheet.max_row, sheet.max_column) != (expected_rows + 1, metadata_width + WAVEFORM_POINTS):
            raise ValueError(f"Unexpected shape in {path.name}: {sheet.max_row}x{sheet.max_column}")
        rows = sheet.iter_rows(values_only=True)
        header = tuple(next(rows))
        if header[:metadata_width] != expected_prefix:
            raise ValueError(f"Unexpected metadata header in {path.name}: {header[:metadata_width]}")
        wave_header = header[metadata_width:]
        if (
            len(wave_header) != WAVEFORM_POINTS
            or not str(wave_header[0]).startswith("0")
            or any(int(float(value)) != index for index, value in enumerate(wave_header[1:], start=1))
        ):
            raise ValueError(f"Unexpected waveform header in {path.name}")

        for excel_row, row in enumerate(rows, start=2):
            if len(row) != metadata_width + WAVEFORM_POINTS:
                raise ValueError(f"Unexpected row width in {path.name}!{excel_row}")
            sample_value, temperature_value, frequency_value, material_value = row[:4]
            sample_float = float(sample_value)
            temperature_float = float(temperature_value)
            frequency = float(frequency_value)
            sample_id = int(sample_float)
            temperature = int(temperature_float)
            material = str(material_value)
            if sample_float != sample_id or temperature_float != temperature:
                raise ValueError(f"Non-integral ID or temperature in {path.name}!{excel_row}")
            if temperature not in TEMPERATURES or material not in MATERIAL_LABELS or not math.isfinite(frequency) or frequency <= 0:
                raise ValueError(f"Invalid metadata in {path.name}!{excel_row}")
            waveform_label = None if metadata_width == 4 else str(row[4])
            if waveform_label is not None and waveform_label not in WAVEFORM_LABELS:
                raise ValueError(f"Invalid waveform label in {path.name}!{excel_row}")
            waveform = np.asarray(row[metadata_width:], dtype=np.float64)
            features, _ = waveform_features(waveform, plateau_fraction)
            record: dict[str, Any] = {
                "sample_id": sample_id,
                "source_excel_row": excel_row,
                "source_file": path.name,
                "source_file_sha256": source_file_hash,
                "waveform_sha256": waveform_sha256(waveform),
                "temperature_C": temperature,
                "temperature_label": str(temperature),
                "frequency_Hz": frequency,
                "log_frequency_Hz": math.log(frequency),
                "material": material,
                **features,
            }
            record["log_b_m_T"] = math.log(float(record["b_m_T"]))
            record["log_b_pp_T"] = math.log(float(record["b_pp_T"]))
            if waveform_label is not None:
                record["waveform"] = waveform_label
            records.append(record)
    finally:
        workbook.close()

    frame = pd.DataFrame.from_records(records).sort_values("sample_id").reset_index(drop=True)
    if list(frame["sample_id"]) != list(range(1, expected_rows + 1)):
        raise ValueError(f"Sample ID coverage failed for {path.name}")
    return frame


def predict_q1(frame: pd.DataFrame, model: Any) -> pd.DataFrame:
    labels = np.asarray(model.predict(frame[SHAPE_FEATURES]), dtype=object)
    probabilities = np.asarray(model.predict_proba(frame[SHAPE_FEATURES]), dtype=np.float64)
    classes = list(model.classes_)
    if set(classes) != set(WAVEFORM_LABELS) or probabilities.shape != (80, 3):
        raise ValueError("Q1 frozen model class schema mismatch")
    if not np.isfinite(probabilities).all() or not np.allclose(probabilities.sum(axis=1), 1.0, atol=1e-12):
        raise ValueError("Q1 frozen model produced invalid probabilities")
    output = frame[
        ["sample_id", "source_excel_row", "source_file_sha256", "waveform_sha256", "temperature_C", "frequency_Hz", "material"]
    ].copy()
    output["predicted_label"] = labels
    output["class_code"] = [CLASS_ENCODING[str(label)] for label in labels]
    names = {"正弦波": "probability_sine", "三角波": "probability_triangle", "梯形波": "probability_trapezoid"}
    for class_index, label in enumerate(classes):
        output[names[label]] = probabilities[:, class_index]
    output["model_id"] = "q1-shape-logistic-full-g4"
    return output


def predict_q4(frame: pd.DataFrame, frozen: dict[str, Any]) -> pd.DataFrame:
    numeric_features = list(frozen["numeric_features"])
    prediction_log, prediction = predict_log_model(
        frozen["model"], float(frozen["smearing_factor"]), frame, numeric_features
    )
    output = frame[
        [
            "sample_id",
            "source_excel_row",
            "source_file_sha256",
            "waveform_sha256",
            "temperature_C",
            "frequency_Hz",
            "material",
            "waveform",
            "b_m_T",
            "b_pp_T",
        ]
    ].copy()
    output["prediction_log"] = prediction_log
    output["predicted_loss_W_per_m3"] = prediction
    output["predicted_loss_rounded_1dp_W_per_m3"] = np.round(prediction, 1)
    output["model_id"] = "q4-hgb-amplitude-and-condition-full-g4"
    return output


def write_submission(q1: pd.DataFrame, q4: pd.DataFrame, destination: Path) -> None:
    workbook = load_workbook(TEMPLATE_FILE)
    try:
        if set(workbook.sheetnames) != {"Sheet1", "Sheet2", "Sheet3"}:
            raise ValueError(f"Unexpected template sheets: {workbook.sheetnames}")
        sheet = workbook["Sheet1"]
        expected_header = ("序号", "附件二（80个样品）励磁波形分类结果", "附件三（400个样品）磁芯损耗预测结果")
        if tuple(cell.value for cell in sheet[1]) != expected_header or (sheet.max_row, sheet.max_column) != (401, 3):
            raise ValueError("Attachment 4 template structure mismatch")
        if any(sheet.cell(row=row, column=2).value is not None for row in range(2, 402)):
            raise ValueError("Attachment 4 Q1 output cells are not blank")
        if any(sheet.cell(row=row, column=3).value is not None for row in range(2, 402)):
            raise ValueError("Attachment 4 Q4 output cells are not blank")
        for item in q1.itertuples(index=False):
            row = int(item.sample_id) + 1
            if int(sheet.cell(row=row, column=1).value) != int(item.sample_id):
                raise ValueError(f"Attachment 4 Q1 ID mismatch at row {row}")
            sheet.cell(row=row, column=2, value=int(item.class_code)).number_format = "0"
        for item in q4.itertuples(index=False):
            row = int(item.sample_id) + 1
            if int(sheet.cell(row=row, column=1).value) != int(item.sample_id):
                raise ValueError(f"Attachment 4 Q4 ID mismatch at row {row}")
            sheet.cell(row=row, column=3, value=float(item.predicted_loss_rounded_1dp_W_per_m3)).number_format = "0.0"
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    finally:
        workbook.close()


def main() -> None:
    args = parse_args()
    config_path = args.config.resolve()
    run = ExperimentRun(
        EXPERIMENT_ID,
        config_path,
        descriptor_path=DESCRIPTOR,
        evidence_root=RUN_ROOT,
        stage="S5",
    )
    try:
        config = load_frozen_config(config_path)
        run.record_contract_hashes(verify_contract_registry(config))
        set_global_seed(int(config["seeds"]["global"]))
        verified_inputs = verify_raw_inputs(config)
        authorization = assert_g4_pass()
        required_inputs = (
            Path(__file__).resolve(),
            Path(__file__).with_name("build_features.py").resolve(),
            Path(__file__).with_name("s3_common.py").resolve(),
            Path(__file__).with_name("s4_common.py").resolve(),
            TEST2_FILE,
            TEST3_FILE,
            TEMPLATE_FILE,
            Q1_MODEL_FILE,
            Q1_METRICS_FILE,
            Q1_RUN_MANIFEST,
            Q4_MODEL_FILE,
            Q4_WINNER_FILE,
            Q4_METRICS_FILE,
            Q4_RUN_MANIFEST,
            G4_REVIEW_FILE,
        )
        for path in required_inputs:
            run.record_input(path)

        q1_manifest = json.loads(Q1_RUN_MANIFEST.read_text(encoding="utf-8"))
        q4_metrics = json.loads(Q4_METRICS_FILE.read_text(encoding="utf-8"))
        q4_winner = json.loads(Q4_WINNER_FILE.read_text(encoding="utf-8"))
        expected_q1_hash = q1_manifest["output_hashes"]["results/raw/s3/q1/logistic_full_model.joblib"]
        if sha256_file(Q1_MODEL_FILE) != expected_q1_hash:
            raise ValueError("Q1 frozen model hash differs from its approved run manifest")
        if sha256_file(Q4_MODEL_FILE) != q4_metrics["final_full_model_sha256"]:
            raise ValueError("Q4 frozen model hash differs from approved ablation metrics")
        if q4_winner.get("feature_variant") != "amplitude_and_condition" or q4_winner.get("winner") != "hgb":
            raise ValueError("Q4 frozen winner differs from the G4-approved model")

        plateau = float(config["feature_contract"]["plateau_relative_slope_threshold"])
        test2 = load_test_features(TEST2_FILE, "test2", plateau)
        test3 = load_test_features(TEST3_FILE, "test3", plateau)
        q1_model = joblib.load(Q1_MODEL_FILE)
        q4_model = joblib.load(Q4_MODEL_FILE)
        q1_predictions = predict_q1(test2, q1_model)
        q4_predictions = predict_q4(test3, q4_model)

        q1_repeat = predict_q1(test2, q1_model)
        q4_repeat = predict_q4(test3, q4_model)
        if not q1_predictions.equals(q1_repeat):
            raise AssertionError("Q1 repeated prediction is not identical")
        if not np.array_equal(
            q4_predictions["predicted_loss_W_per_m3"].to_numpy(),
            q4_repeat["predicted_loss_W_per_m3"].to_numpy(),
        ):
            raise AssertionError("Q4 repeated prediction is not bitwise identical")

        prediction_dir = RAW_S5_ROOT / "predictions"
        submission_path = RAW_S5_ROOT / "submission" / "附件四（Excel表）.xlsx"
        q1_path = prediction_dir / "q1_attachment2_predictions.csv"
        q4_path = prediction_dir / "q4_attachment3_predictions.csv"
        write_csv(q1_path, q1_predictions)
        write_csv(q4_path, q4_predictions)
        write_submission(q1_predictions, q4_predictions, submission_path)

        template_hash_after = sha256_file(TEMPLATE_FILE)
        if template_hash_after != verified_inputs["src/附件四（Excel表）.xlsx"]:
            raise AssertionError("Original attachment 4 changed during output generation")
        for relative, digest in verified_inputs.items():
            if sha256_file(PROJECT_ROOT / relative) != digest:
                raise AssertionError(f"Original input changed during output generation: {relative}")

        q1_counts = Counter(q1_predictions["predicted_label"])
        q1_selected_ids = [1, 5, 15, 25, 35, 45, 55, 65, 75, 80]
        q4_selected_ids = [16, 76, 98, 126, 168, 230, 271, 338, 348, 379]
        summary = {
            "status": "PASS",
            "prediction_policy": "one_time_frozen_prediction_after_G4_PASS",
            "selection_or_tuning_with_test_attachments": False,
            "g4_authorization": authorization,
            "q1": {
                "model_id": "q1-shape-logistic-full-g4",
                "row_count": int(len(q1_predictions)),
                "class_counts": {label: int(q1_counts[label]) for label in WAVEFORM_LABELS},
                "specified_samples": q1_predictions[q1_predictions["sample_id"].isin(q1_selected_ids)][
                    ["sample_id", "predicted_label", "class_code"]
                ].to_dict(orient="records"),
                "repeated_prediction_identical": True,
            },
            "q4": {
                "model_id": "q4-hgb-amplitude-and-condition-full-g4",
                "feature_variant": q4_winner["feature_variant"],
                "numeric_feature_count": len(q4_model["numeric_features"]),
                "row_count": int(len(q4_predictions)),
                "prediction_min_W_per_m3": float(q4_predictions["predicted_loss_W_per_m3"].min()),
                "prediction_max_W_per_m3": float(q4_predictions["predicted_loss_W_per_m3"].max()),
                "specified_samples": q4_predictions[q4_predictions["sample_id"].isin(q4_selected_ids)][
                    ["sample_id", "predicted_loss_rounded_1dp_W_per_m3"]
                ].to_dict(orient="records"),
                "repeated_prediction_bitwise_identical": True,
                "unknown_ground_truth": True,
            },
            "test_domain": {
                "attachment2_frequency_range_Hz": [float(test2["frequency_Hz"].min()), float(test2["frequency_Hz"].max())],
                "attachment3_frequency_range_Hz": [float(test3["frequency_Hz"].min()), float(test3["frequency_Hz"].max())],
                "attachment3_materials": sorted(test3["material"].unique()),
                "attachment3_temperatures_C": sorted(int(value) for value in test3["temperature_C"].unique()),
                "attachment3_waveforms": sorted(test3["waveform"].unique()),
            },
            "generated_at": datetime.now(TZ).isoformat(),
        }
        summary_path = RAW_S5_ROOT / "prediction_summary.json"
        write_json(summary_path, summary)

        q4_run = json.loads(Q4_RUN_MANIFEST.read_text(encoding="utf-8"))
        freeze = {
            "freeze_id": "S5-FREEZE-2024C-V1",
            "status": "FROZEN_PENDING_G5",
            "g4_authorization": authorization,
            "feature_contract": "wave-v1",
            "q1": {
                "model": "shape-only multinomial logistic regression",
                "model_path": "results/raw/s3/q1/logistic_full_model.joblib",
                "model_sha256": sha256_file(Q1_MODEL_FILE),
                "source_implementation_commit": q1_manifest["git_commit"],
                "class_encoding": CLASS_ENCODING,
                "feature_names": SHAPE_FEATURES,
            },
            "q4": {
                "model": "HistGradientBoostingRegressor on log loss",
                "model_path": "results/raw/main/q4/final_full_model.joblib",
                "model_sha256": sha256_file(Q4_MODEL_FILE),
                "source_implementation_commit": q4_run["git_commit"],
                "feature_variant": q4_winner["feature_variant"],
                "numeric_features": q4_model["numeric_features"],
                "categorical_features": ["material", "waveform", "temperature_label"],
                "parameters": q4_model["parameters"],
                "smearing_factor": float(q4_model["smearing_factor"]),
            },
            "official_input_hashes": verified_inputs,
            "raw_prediction_outputs": {
                "q1_csv": {"path": "results/raw/s5/predictions/q1_attachment2_predictions.csv", "sha256": sha256_file(q1_path)},
                "q4_csv": {"path": "results/raw/s5/predictions/q4_attachment3_predictions.csv", "sha256": sha256_file(q4_path)},
                "attachment4": {"path": "results/raw/s5/submission/附件四（Excel表）.xlsx", "sha256": sha256_file(submission_path)},
            },
            "limitations": [
                "Attachment 2 has no disclosed ground truth; predictions are not accuracy evidence.",
                "Attachment 3 has no disclosed core-loss ground truth; predictions are not external-test metrics.",
                "Q4 leave-one-level stress limits remain applicable.",
                "Q5 unique recommendation remains unauthorized because the primary observed-region Jaccard gate failed.",
            ],
        }
        freeze_path = RAW_S5_ROOT / "model_freeze.json"
        write_json(freeze_path, freeze)

        for path in (q1_path, q4_path, submission_path, summary_path, freeze_path):
            run.record_output(path)
        run.log(json.dumps({"q1_counts": summary["q1"]["class_counts"], "q4_range": [summary["q4"]["prediction_min_W_per_m3"], summary["q4"]["prediction_max_W_per_m3"]]}, ensure_ascii=False, sort_keys=True))
        run.finish("PASS")
    except BaseException as exc:
        run.fail(exc)
        raise


if __name__ == "__main__":
    main()
