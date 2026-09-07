"""Build the deterministic wave-v1 feature table and leakage-safe outer folds."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import struct
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.model_selection import GroupKFold, StratifiedGroupKFold

from s3_common import (
    CLASS_ENCODING,
    DEFAULT_CONFIG,
    MATERIAL_LABELS,
    PROJECT_ROOT,
    Q4_CATEGORICAL_FEATURES,
    Q4_NUMERIC_FEATURES,
    S3_RESULTS_ROOT,
    SHAPE_FEATURES,
    TEMPERATURES,
    WAVEFORM_LABELS,
    ExperimentRun,
    load_frozen_config,
    set_global_seed,
    sha256_file,
    verify_contract_registry,
    verify_raw_inputs,
    write_csv,
    write_json,
)


EXPERIMENT_ID = "EXP-S3-DATA-001"
TRAIN_FILE = "附件一（训练集）.xlsx"
TRAIN_SHEETS = {"材料1": 3400, "材料2": 3000, "材料3": 3200, "材料4": 2800}
WAVEFORM_POINTS = 1024


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    return parser.parse_args()


def waveform_sha256(waveform: np.ndarray) -> str:
    canonical = np.asarray(waveform, dtype="<f8").copy()
    canonical[canonical == 0.0] = 0.0
    return hashlib.sha256(canonical.tobytes(order="C")).hexdigest().upper()


def exact_record_sha256(
    material: str,
    temperature: float,
    frequency: float,
    loss: float,
    waveform_label: str,
    waveform: np.ndarray,
) -> str:
    digest = hashlib.sha256()
    for text in (material, waveform_label):
        encoded = text.encode("utf-8")
        digest.update(len(encoded).to_bytes(4, "little"))
        digest.update(encoded)
    digest.update(struct.pack("<ddd", temperature, frequency, loss))
    canonical = np.asarray(waveform, dtype="<f8").copy()
    canonical[canonical == 0.0] = 0.0
    digest.update(canonical.tobytes(order="C"))
    return digest.hexdigest().upper()


def shape_signature(block_means: np.ndarray, step: float) -> str:
    quantized = np.rint(block_means / step).astype("<i4")
    digest = hashlib.sha256()
    digest.update(b"near-shape-v1\0")
    digest.update(struct.pack("<d", step))
    digest.update(quantized.tobytes(order="C"))
    return digest.hexdigest().upper()


def waveform_features(waveform: np.ndarray, plateau_fraction: float) -> tuple[dict[str, float], np.ndarray]:
    values = np.asarray(waveform, dtype=np.float64)
    if values.shape != (WAVEFORM_POINTS,) or not np.isfinite(values).all():
        raise ValueError("Waveform must contain exactly 1024 finite values")

    mean = float(np.mean(values))
    b_min = float(np.min(values))
    b_max = float(np.max(values))
    b_m = float(np.max(np.abs(values)))
    b_pp = b_max - b_min
    if not math.isfinite(b_pp) or b_pp <= 0 or b_m <= 0:
        raise ValueError("Flat or invalid waveform")
    b_rms = float(np.sqrt(np.mean(values**2)))
    b_std = float(np.std(values))
    if b_rms <= 0 or b_std <= 0:
        raise ValueError("Invalid RMS or standard deviation")

    normalized = (values - mean) / b_pp
    norm_std = float(np.std(normalized))
    standardized = (normalized - float(np.mean(normalized))) / norm_std
    diff = np.roll(values, -1) - values
    shape_diff = np.roll(normalized, -1) - normalized
    abs_diff = np.abs(diff)
    abs_shape_diff = np.abs(shape_diff)
    max_abs_diff = float(np.max(abs_diff))
    if max_abs_diff <= 0:
        raise ValueError("Waveform has no finite periodic slope")
    plateau_threshold = plateau_fraction * max_abs_diff
    significant_signs = np.sign(diff[abs_diff > plateau_threshold])
    if significant_signs.size < 2:
        sign_changes = 0
    else:
        sign_changes = int(np.sum(significant_signs != np.roll(significant_signs, 1)))

    spectrum = np.fft.rfft(normalized - np.mean(normalized))
    energy = np.abs(spectrum) ** 2
    positive_energy = float(np.sum(energy[1:]))
    fundamental_energy = float(energy[1])
    tolerance = np.finfo(np.float64).eps * max(1.0, positive_energy)
    if positive_energy <= tolerance or fundamental_energy <= tolerance:
        raise ValueError("Fundamental FFT energy is below machine tolerance")
    harmonic_ratios = energy[1:11] / positive_energy

    quantiles = np.quantile(values, [0.05, 0.25, 0.5, 0.75, 0.95])
    shape_quantiles = np.quantile(normalized, [0.05, 0.25, 0.5, 0.75, 0.95])
    diff_quantiles = np.quantile(abs_diff, [0.5, 0.9, 0.95])
    shape_diff_quantiles = np.quantile(abs_shape_diff, [0.5, 0.9, 0.95])

    features: dict[str, float] = {
        "b_mean_T": mean,
        "b_m_T": b_m,
        "b_pp_T": b_pp,
        "b_half_pp_T": b_pp / 2.0,
        "b_rms_T": b_rms,
        "b_crest_factor": b_m / b_rms,
        "b_std_T": b_std,
        "b_q05_T": float(quantiles[0]),
        "b_q25_T": float(quantiles[1]),
        "b_median_T": float(quantiles[2]),
        "b_q75_T": float(quantiles[3]),
        "b_q95_T": float(quantiles[4]),
        "diff_max_abs_T": max_abs_diff,
        "diff_mean_abs_T": float(np.mean(abs_diff)),
        "diff_q50_abs_T": float(diff_quantiles[0]),
        "diff_q90_abs_T": float(diff_quantiles[1]),
        "diff_q95_abs_T": float(diff_quantiles[2]),
        "total_variation_T": float(np.sum(abs_diff)),
        "shape_std": norm_std,
        "shape_q05": float(shape_quantiles[0]),
        "shape_q25": float(shape_quantiles[1]),
        "shape_median": float(shape_quantiles[2]),
        "shape_q75": float(shape_quantiles[3]),
        "shape_q95": float(shape_quantiles[4]),
        "shape_skew": float(np.mean(standardized**3)),
        "shape_kurtosis_excess": float(np.mean(standardized**4) - 3.0),
        "shape_crest_factor": float(np.max(np.abs(normalized)) / np.sqrt(np.mean(normalized**2))),
        "shape_diff_max_abs": float(np.max(abs_shape_diff)),
        "shape_diff_mean_abs": float(np.mean(abs_shape_diff)),
        "shape_diff_q50_abs": float(shape_diff_quantiles[0]),
        "shape_diff_q90_abs": float(shape_diff_quantiles[1]),
        "shape_diff_q95_abs": float(shape_diff_quantiles[2]),
        "shape_total_variation": float(np.sum(abs_shape_diff)),
        "shape_positive_slope_ratio": float(np.mean(shape_diff > 0)),
        "shape_negative_slope_ratio": float(np.mean(shape_diff < 0)),
        "shape_plateau_ratio": float(np.mean(abs_diff <= plateau_threshold)),
        "shape_slope_sign_changes": float(sign_changes),
        "shape_fft_harmonic_to_fundamental": float(np.sum(energy[2:11]) / fundamental_energy),
    }
    for index, ratio in enumerate(harmonic_ratios, start=1):
        features[f"shape_fft_h{index}_energy_ratio"] = float(ratio)

    aligned = np.roll(normalized, -int(np.argmax(normalized)))
    block_means = aligned.reshape(64, 16).mean(axis=1)
    return features, block_means


def choose_near_shape_groups(frame: pd.DataFrame, block_matrix: np.ndarray, steps: list[float]) -> tuple[list[str], float, dict[str, int]]:
    for step in steps:
        signatures = [shape_signature(row, step) for row in block_matrix]
        group_counts = {
            label: len({signature for signature, actual in zip(signatures, frame["waveform"], strict=True) if actual == label})
            for label in WAVEFORM_LABELS
        }
        if min(group_counts.values()) >= 5:
            return [f"near-shape-v1:{step:g}:{signature}" for signature in signatures], step, group_counts
    exact = [f"exact-wave:{value}" for value in frame["waveform_sha256"]]
    group_counts = {
        label: frame.loc[frame["waveform"] == label, "waveform_sha256"].nunique()
        for label in WAVEFORM_LABELS
    }
    if min(group_counts.values()) < 5:
        raise ValueError(f"Insufficient distinct waveform groups after fallback: {group_counts}")
    return exact, math.nan, group_counts


def choose_q1_folds(frame: pd.DataFrame, seed: int, fold_candidates: list[int]) -> tuple[np.ndarray, int]:
    labels = frame["waveform"].to_numpy()
    groups = frame["near_shape_group"].to_numpy()
    required = set(WAVEFORM_LABELS)
    for folds in fold_candidates:
        splitter = StratifiedGroupKFold(n_splits=folds, shuffle=True, random_state=seed)
        assignment = np.full(len(frame), -1, dtype=np.int16)
        valid = True
        for fold, (train_index, valid_index) in enumerate(splitter.split(frame, labels, groups)):
            if set(labels[valid_index]) != required or set(labels[train_index]) != required:
                valid = False
                break
            assignment[valid_index] = fold
        if valid and np.all(assignment >= 0):
            return assignment, folds
    raise ValueError("Q1 grouped folds cannot preserve all waveform classes with 5/4/3 folds")


def choose_regression_folds(frame: pd.DataFrame, seed: int, fold_candidates: list[int]) -> tuple[np.ndarray, int]:
    groups = frame["condition_group"].to_numpy()
    q2_mask = (frame["material"] == "材料1") & (frame["waveform"] == "正弦波")
    required_materials = set(MATERIAL_LABELS)
    required_waveforms = set(WAVEFORM_LABELS)
    required_temperatures = set(TEMPERATURES)
    for folds in fold_candidates:
        splitter = GroupKFold(n_splits=folds, shuffle=True, random_state=seed)
        assignment = np.full(len(frame), -1, dtype=np.int16)
        valid = True
        for fold, (train_index, valid_index) in enumerate(splitter.split(frame, groups=groups)):
            validation = frame.iloc[valid_index]
            training = frame.iloc[train_index]
            if (
                set(validation["material"]) != required_materials
                or set(validation["waveform"]) != required_waveforms
                or set(validation["temperature_C"].astype(int)) != required_temperatures
                or set(training["material"]) != required_materials
                or set(training["waveform"]) != required_waveforms
                or set(training["temperature_C"].astype(int)) != required_temperatures
            ):
                valid = False
                break
            q2_validation = validation[(validation["material"] == "材料1") & (validation["waveform"] == "正弦波")]
            q2_training = training[(training["material"] == "材料1") & (training["waveform"] == "正弦波")]
            if set(q2_validation["temperature_C"].astype(int)) != required_temperatures or set(q2_training["temperature_C"].astype(int)) != required_temperatures:
                valid = False
                break
            assignment[valid_index] = fold
        if valid and np.all(assignment >= 0) and int(np.sum(q2_mask)) == 1067:
            return assignment, folds
    raise ValueError("Regression grouped folds cannot preserve required factor coverage with 5/4/3 folds")


def validate_fold_lineage(frame: pd.DataFrame) -> dict[str, Any]:
    checks: dict[str, Any] = {}
    for group_column, fold_column in (
        ("near_shape_group", "q1_outer_fold"),
        ("condition_group", "regression_outer_fold"),
        ("exact_record_sha256", "q1_outer_fold"),
        ("exact_record_sha256", "regression_outer_fold"),
    ):
        maximum = int(frame.groupby(group_column, observed=True)[fold_column].nunique().max())
        if maximum != 1:
            raise AssertionError(f"{group_column} crosses {fold_column}")
        checks[f"{group_column}_max_{fold_column}_count"] = maximum
    return checks


def load_training_frame(config: dict[str, Any], run: ExperimentRun) -> tuple[pd.DataFrame, np.ndarray]:
    source_path = PROJECT_ROOT / "src" / TRAIN_FILE
    run.record_input(source_path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    blocks: list[np.ndarray] = []
    plateau_fraction = float(config["feature_contract"]["plateau_relative_slope_threshold"])
    try:
        if workbook.sheetnames != list(TRAIN_SHEETS):
            raise ValueError(f"Unexpected training sheets: {workbook.sheetnames}")
        for material, expected_rows in TRAIN_SHEETS.items():
            sheet = workbook[material]
            rows = sheet.iter_rows(values_only=True)
            header = tuple(next(rows))
            if header[:4] != ("温度，oC", "频率，Hz", "磁芯损耗，w/m3", "励磁波形"):
                raise ValueError(f"Unexpected header in {material}: {header[:4]}")
            count = 0
            for excel_row, row in enumerate(rows, start=2):
                count += 1
                temperature, frequency, loss, waveform_label = row[:4]
                waveform = np.asarray(row[4:], dtype=np.float64)
                if waveform_label not in WAVEFORM_LABELS:
                    raise ValueError(f"Unknown waveform label {waveform_label!r} at {material}!{excel_row}")
                values = [float(temperature), float(frequency), float(loss)]
                if not np.isfinite(values).all() or values[0] not in TEMPERATURES or values[1] <= 0 or values[2] <= 0:
                    raise ValueError(f"Invalid metadata at {material}!{excel_row}")
                features, block_means = waveform_features(waveform, plateau_fraction)
                wave_hash = waveform_sha256(waveform)
                row_id = f"{TRAIN_FILE}|{material}|{excel_row}"
                record: dict[str, Any] = {
                    "row_id": row_id,
                    "source_file": TRAIN_FILE,
                    "source_sheet": material,
                    "source_excel_row": excel_row,
                    "source_file_sha256": config["input_hashes"][f"src/{TRAIN_FILE}"],
                    "material": material,
                    "temperature_C": int(values[0]),
                    "temperature_label": str(int(values[0])),
                    "frequency_Hz": values[1],
                    "core_loss_W_per_m3": values[2],
                    "waveform": str(waveform_label),
                    "waveform_code": CLASS_ENCODING[str(waveform_label)],
                    "waveform_sha256": wave_hash,
                    "exact_record_sha256": exact_record_sha256(
                        material,
                        values[0],
                        values[1],
                        values[2],
                        str(waveform_label),
                        waveform,
                    ),
                    "log_frequency_Hz": math.log(values[1]),
                    **features,
                }
                record["log_b_m_T"] = math.log(record["b_m_T"])
                record["log_b_pp_T"] = math.log(record["b_pp_T"])
                gf = math.floor(math.log10(values[1]) / float(config["validation_contract"]["log10_frequency_bin_width"]))
                gb = math.floor(math.log10(record["b_m_T"]) / float(config["validation_contract"]["log10_b_peak_bin_width"]))
                record["condition_group"] = f"{material}|{int(values[0])}|{waveform_label}|{gf}|{gb}"
                record["energy_proxy_Hz_T"] = values[1] * record["b_m_T"]
                records.append(record)
                blocks.append(block_means)
            if count != expected_rows:
                raise ValueError(f"Expected {expected_rows} rows in {material}, got {count}")
            run.log(f"loaded sheet={material} rows={count}")
    finally:
        workbook.close()
    frame = pd.DataFrame.from_records(records)
    block_matrix = np.vstack(blocks)
    if len(frame) != 12400 or block_matrix.shape != (12400, 64):
        raise ValueError(f"Unexpected training shape frame={frame.shape}, blocks={block_matrix.shape}")
    if frame["row_id"].duplicated().any():
        raise ValueError("row_id is not unique")
    numeric = frame.select_dtypes(include=[np.number])
    if not np.isfinite(numeric.to_numpy(dtype=np.float64)).all():
        raise ValueError("Feature table contains non-finite numeric values")
    return frame, block_matrix


def main() -> None:
    args = parse_args()
    config_path = args.config.resolve()
    run = ExperimentRun(EXPERIMENT_ID, config_path, manifest_aliases=[S3_RESULTS_ROOT / "run_manifest.json"])
    try:
        config = load_frozen_config(config_path)
        seed = int(config["seeds"]["global"])
        set_global_seed(seed)
        run.record_input(Path(__file__).resolve())
        run.record_input(Path(__file__).with_name("s3_common.py").resolve())
        contract_hashes = verify_contract_registry(config)
        inputs_before = verify_raw_inputs(config)
        for relative in config["input_hashes"]:
            usage_label = relative if relative == f"src/{TRAIN_FILE}" else f"integrity_hash_only:{relative}"
            run.record_input(PROJECT_ROOT / relative, label=usage_label)
        run.log(f"verified contracts={len(contract_hashes)} raw_inputs={len(inputs_before)}")

        frame, block_matrix = load_training_frame(config, run)
        quantization_steps = [float(value) for value in config["validation_contract"]["near_shape_quantization"]]
        near_groups, selected_step, near_counts = choose_near_shape_groups(frame, block_matrix, quantization_steps)
        frame["near_shape_group"] = near_groups
        frame["near_shape_quantization"] = selected_step

        fold_candidates = [int(config["validation_contract"]["outer_folds_target"]), *map(int, config["validation_contract"]["outer_folds_fallback"])]
        frame["q1_outer_fold"], q1_folds = choose_q1_folds(frame, seed, fold_candidates)
        frame["regression_outer_fold"], regression_folds = choose_regression_folds(frame, seed, fold_candidates)
        lineage_checks = validate_fold_lineage(frame)

        duplicate_sizes = frame.groupby("exact_record_sha256", observed=True).size()
        duplicate_extra_rows = int((duplicate_sizes - 1).clip(lower=0).sum())
        if duplicate_extra_rows != 1:
            raise AssertionError(f"Expected one extra exact duplicate row, got {duplicate_extra_rows}")

        feature_path = S3_RESULTS_ROOT / "data" / "feature_table.csv"
        fold_path = S3_RESULTS_ROOT / "folds" / "fold_assignments.csv"
        schema_path = S3_RESULTS_ROOT / "data" / "feature_schema.json"
        metrics_path = run.output_dir / "metrics.json"
        canonical_metrics_path = S3_RESULTS_ROOT / "data" / "build_metrics.json"
        rebuild_paths = {
            "feature_table_sha256": feature_path,
            "fold_assignments_sha256": fold_path,
            "feature_schema_sha256": schema_path,
        }
        prior_output_hashes = {
            label: sha256_file(path)
            for label, path in rebuild_paths.items()
            if path.is_file()
        }
        if prior_output_hashes and len(prior_output_hashes) != len(rebuild_paths):
            raise AssertionError("Only part of the deterministic data bundle existed before rebuild")

        fold_columns = [
            "row_id",
            "near_shape_group",
            "condition_group",
            "exact_record_sha256",
            "q1_outer_fold",
            "regression_outer_fold",
        ]
        write_csv(feature_path, frame.drop(columns=["q1_outer_fold", "regression_outer_fold"]))
        write_csv(fold_path, frame[fold_columns])

        schema = {
            "schema_version": "wave-v1-s3-1.0",
            "rows": len(frame),
            "waveform_points": WAVEFORM_POINTS,
            "primary_b_peak": "max(abs(B))",
            "waveform_sha256": "float64 little-endian; negative zero normalized to positive zero",
            "near_shape_quantization": None if math.isnan(selected_step) else selected_step,
            "q1_shape_features": SHAPE_FEATURES,
            "q4_numeric_features": Q4_NUMERIC_FEATURES,
            "q4_categorical_features": Q4_CATEGORICAL_FEATURES,
            "class_encoding": CLASS_ENCODING,
            "source_file_sha256": inputs_before[f"src/{TRAIN_FILE}"],
            "contract_hashes": contract_hashes,
        }
        write_json(schema_path, schema)
        current_output_hashes = {label: sha256_file(path) for label, path in rebuild_paths.items()}
        rebuild_match = None if not prior_output_hashes else prior_output_hashes == current_output_hashes
        if rebuild_match is False:
            raise AssertionError(
                f"Deterministic rebuild hash mismatch: before={prior_output_hashes}; after={current_output_hashes}"
            )
        metrics = {
            "status": "PASS",
            "rows": len(frame),
            "numeric_feature_count": len(Q4_NUMERIC_FEATURES),
            "q1_shape_feature_count": len(SHAPE_FEATURES),
            "q1_outer_folds": q1_folds,
            "regression_outer_folds": regression_folds,
            "near_shape_quantization": None if math.isnan(selected_step) else selected_step,
            "near_shape_distinct_by_class": near_counts,
            "condition_group_count": int(frame["condition_group"].nunique()),
            "exact_duplicate_extra_rows": duplicate_extra_rows,
            "lineage_checks": lineage_checks,
            "rebuild_prior_bundle_present": bool(prior_output_hashes),
            "rebuild_hashes_match": rebuild_match,
            "raw_input_usage": {
                f"src/{TRAIN_FILE}": "feature_source",
                "all_other_registered_inputs": "integrity_hash_only_not_loaded_for_modeling",
            },
            **current_output_hashes,
        }
        write_json(metrics_path, metrics)
        write_json(canonical_metrics_path, metrics)

        inputs_after = verify_raw_inputs(config)
        if inputs_before != inputs_after:
            raise AssertionError("Raw input hashes changed during feature build")
        for path in (feature_path, fold_path, schema_path, metrics_path, canonical_metrics_path):
            run.record_output(path)
        run.log(json.dumps(metrics, ensure_ascii=False, sort_keys=True))
        run.finish("PASS")
    except BaseException as exc:
        run.fail(exc)
        raise


if __name__ == "__main__":
    main()
