"""Verify S5 frozen outputs and promote only an explicit whitelist to verified."""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from run_s5_predictions import (
    G4_REVIEWED_COMMIT,
    Q1_MODEL_FILE,
    Q1_RUN_MANIFEST,
    Q4_METRICS_FILE,
    Q4_MODEL_FILE,
    RAW_S5_ROOT,
    TEMPLATE_FILE,
    TEST2_FILE,
    TEST3_FILE,
    assert_g4_pass,
    load_test_features,
    predict_q1,
    predict_q4,
)
from s3_common import (
    CLASS_ENCODING,
    DEFAULT_CONFIG,
    PROJECT_ROOT,
    ExperimentRun,
    load_frozen_config,
    sha256_file,
    verify_contract_registry,
    verify_raw_inputs,
    write_json,
)


EXPERIMENT_ID = "EXP-S5-VERIFY-001"
DESCRIPTOR = PROJECT_ROOT / "experiments" / "main" / f"{EXPERIMENT_ID}.json"
RUN_ROOT = RAW_S5_ROOT / "runs"
VERIFIED_ROOT = PROJECT_ROOT / "results" / "verified"
RAW_Q1_PREDICTIONS = RAW_S5_ROOT / "predictions" / "q1_attachment2_predictions.csv"
RAW_Q4_PREDICTIONS = RAW_S5_ROOT / "predictions" / "q4_attachment3_predictions.csv"
RAW_SUBMISSION = RAW_S5_ROOT / "submission" / "附件四（Excel表）.xlsx"
RAW_FREEZE = RAW_S5_ROOT / "model_freeze.json"
RAW_SUMMARY = RAW_S5_ROOT / "prediction_summary.json"
S4_VERIFICATION = PROJECT_ROOT / "results" / "raw" / "main" / "verification_report.json"


PROMOTION_MAP = {
    "results/raw/s1/dataset_summary.csv": "data/dataset_summary.csv",
    "results/raw/s1/quality_checks.csv": "data/quality_checks.csv",
    "results/raw/s1/numeric_ranges.csv": "data/numeric_ranges.csv",
    "results/raw/s1/categorical_counts.csv": "data/categorical_counts.csv",
    "results/raw/s1/data_profile.json": "data/data_profile.json",
    "results/raw/s3/q1/baseline_metrics.json": "q1/oof_metrics.json",
    "results/raw/s3/q1/confusion_matrix.csv": "q1/confusion_matrix.csv",
    "results/raw/main/q1/invariance_metrics.csv": "q1/invariance_metrics.csv",
    "results/raw/main/q1/leave_one_material_out.csv": "q1/leave_one_material_out.csv",
    "results/raw/main/q1/feature_auxiliary_ablation.csv": "q1/feature_auxiliary_ablation.csv",
    "results/raw/s5/predictions/q1_attachment2_predictions.csv": "q1/attachment2_predictions.csv",
    "results/raw/main/q2/quadratic_temperature_metrics.json": "q2/quadratic_temperature_metrics.json",
    "results/raw/main/q2/quadratic_temperature_parameters.csv": "q2/quadratic_temperature_parameters.csv",
    "results/raw/main/q2/quadratic_temperature_by_temperature.csv": "q2/quadratic_temperature_by_temperature.csv",
    "results/raw/main/q2/leave_one_temperature_comparison.csv": "q2/leave_one_temperature_comparison.csv",
    "results/raw/main/q2/sensitivity_metrics.json": "q2/sensitivity_metrics.json",
    "results/raw/main/q2/sensitivity_summary.csv": "q2/sensitivity_summary.csv",
    "results/raw/main/q3/interaction_metrics.json": "q3/interaction_metrics.json",
    "results/raw/main/q3/bootstrap_metrics.json": "q3/bootstrap_metrics.json",
    "results/raw/main/q3/bootstrap_pairwise_interaction_intervals.csv": "q3/bootstrap_pairwise_interaction_intervals.csv",
    "results/raw/main/q3/bootstrap_main_effect_intervals.csv": "q3/bootstrap_main_effect_intervals.csv",
    "results/raw/main/q3/bootstrap_adjusted_cell_intervals.csv": "q3/bootstrap_adjusted_cell_intervals.csv",
    "results/raw/main/q3/bootstrap_lowest_combination_frequency.csv": "q3/bootstrap_lowest_combination_frequency.csv",
    "results/raw/main/q3/sensitivity_metrics.json": "q3/sensitivity_metrics.json",
    "results/raw/main/q3/sensitivity_summary.csv": "q3/sensitivity_summary.csv",
    "results/raw/main/q4/hgb_metrics.json": "q4/hgb_metrics.json",
    "results/raw/main/q4/ablation_metrics.json": "q4/ablation_metrics.json",
    "results/raw/main/q4/ablation_metrics.csv": "q4/ablation_metrics.csv",
    "results/raw/main/q4/final_oof_predictions.csv": "q4/final_oof_predictions.csv",
    "results/raw/main/q4/hgb_subgroup_metrics.csv": "q4/hgb_subgroup_metrics.csv",
    "results/raw/main/q4/stress_subset_metrics.csv": "q4/stress_subset_metrics.csv",
    "results/raw/main/q4/leave_one_level_out_metrics.csv": "q4/leave_one_level_out_metrics.csv",
    "results/raw/main/q4/stress_metrics.json": "q4/stress_metrics.json",
    "results/raw/s5/predictions/q4_attachment3_predictions.csv": "q4/attachment3_predictions.csv",
    "results/raw/main/q5/robustness_metrics.json": "q5/robustness_metrics.json",
    "results/raw/main/q5/sensitivity_metrics.json": "q5/sensitivity_metrics.json",
    "results/raw/main/q5/sensitivity_scenarios.csv": "q5/sensitivity_scenarios.csv",
    "results/raw/main/q5/final_oof_pareto.csv": "q5/final_oof_pareto.csv",
    "results/raw/main/q5/final_full_fit_reference_pareto.csv": "q5/final_full_fit_reference_pareto.csv",
    "results/raw/main/q5/bootstrap_region_stability.csv": "q5/bootstrap_region_stability.csv",
    "results/raw/main/q5/bootstrap_candidate_stability.csv": "q5/bootstrap_candidate_stability.csv",
    "results/raw/main/q5/heldout_fold_region_support.csv": "q5/heldout_fold_region_support.csv",
    "results/raw/main/q5/representative_conditions.json": "q5/representative_conditions.json",
    "results/raw/main/q5/conflict_thresholds.csv": "q5/conflict_thresholds.csv",
    "results/raw/s5/submission/附件四（Excel表）.xlsx": "submission/附件四（Excel表）.xlsx",
    "results/raw/s5/model_freeze.json": "model_freeze.json",
    "results/raw/s5/prediction_summary.json": "prediction_summary.json",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    return parser.parse_args()


def compare_numeric(
    name: str,
    actual: pd.Series,
    expected: pd.Series,
    *,
    atol: float,
    checks: list[dict[str, Any]],
    rtol: float = 0.0,
) -> None:
    left = actual.to_numpy(dtype=np.float64)
    right = expected.to_numpy(dtype=np.float64)
    passed = left.shape == right.shape and np.isfinite(left).all() and np.allclose(left, right, rtol=rtol, atol=atol)
    maximum = float(np.max(np.abs(left - right))) if left.shape == right.shape and left.size else None
    record_check(checks, name, passed, {"maximum_absolute_difference": maximum, "atol": atol, "rtol": rtol})


def record_check(checks: list[dict[str, Any]], name: str, passed: bool, detail: Any) -> None:
    checks.append({"check": name, "status": "PASS" if passed else "FAIL", "detail": detail})
    if not passed:
        raise AssertionError(f"{name}: {detail}")


def verify_submission(q1: pd.DataFrame, q4: pd.DataFrame, checks: list[dict[str, Any]]) -> None:
    workbook = load_workbook(RAW_SUBMISSION, read_only=True, data_only=True)
    try:
        record_check(checks, "submission_sheet_set", set(workbook.sheetnames) == {"Sheet1", "Sheet2", "Sheet3"}, workbook.sheetnames)
        sheet = workbook["Sheet1"]
        record_check(checks, "submission_shape", (sheet.max_row, sheet.max_column) == (401, 3), [sheet.max_row, sheet.max_column])
        rows = list(sheet.iter_rows(values_only=True))
        expected_header = ("序号", "附件二（80个样品）励磁波形分类结果", "附件三（400个样品）磁芯损耗预测结果")
        record_check(checks, "submission_header", tuple(rows[0]) == expected_header, list(rows[0]))
        record_check(checks, "submission_ids", [int(row[0]) for row in rows[1:]] == list(range(1, 401)), "1..400")
        q1_values = [int(rows[index][1]) for index in range(1, 81)]
        record_check(checks, "submission_q1_mapping", q1_values == list(q1["class_code"].astype(int)), {"rows": len(q1_values)})
        record_check(checks, "submission_q1_unused_rows_blank", all(rows[index][1] is None for index in range(81, 401)), "sample_ids=81..400")
        q4_values = np.asarray([float(rows[index][2]) for index in range(1, 401)], dtype=np.float64)
        expected_q4 = q4["predicted_loss_rounded_1dp_W_per_m3"].to_numpy(dtype=np.float64)
        record_check(checks, "submission_q4_mapping", np.array_equal(q4_values, expected_q4), {"rows": len(q4_values)})
    finally:
        workbook.close()


def promote_whitelist(checks: list[dict[str, Any]]) -> dict[str, dict[str, str]]:
    allowed = {Path(value).as_posix() for value in PROMOTION_MAP.values()}
    allowed.update({"result_registry.md", "verification_report.json", "provenance.json"})
    if VERIFIED_ROOT.exists():
        unexpected = sorted(
            path.relative_to(VERIFIED_ROOT).as_posix()
            for path in VERIFIED_ROOT.rglob("*")
            if path.is_file() and path.relative_to(VERIFIED_ROOT).as_posix() not in allowed
        )
    else:
        unexpected = []
    record_check(checks, "verified_contains_no_unmanaged_files", not unexpected, unexpected)

    provenance: dict[str, dict[str, str]] = {}
    for source_relative, destination_relative in PROMOTION_MAP.items():
        source = PROJECT_ROOT / source_relative
        destination = VERIFIED_ROOT / destination_relative
        if not source.is_file():
            raise FileNotFoundError(f"Promotion source missing: {source_relative}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
        source_hash = sha256_file(source)
        destination_hash = sha256_file(destination)
        record_check(checks, f"promoted_hash::{destination_relative}", source_hash == destination_hash, destination_hash)
        provenance[destination_relative] = {
            "source": source_relative,
            "source_sha256": source_hash,
            "verified_sha256": destination_hash,
        }
    return provenance


def build_registry(summary: dict[str, Any]) -> str:
    q1_samples = ", ".join(f"{row['sample_id']}→{row['class_code']}" for row in summary["q1"]["specified_samples"])
    q4_samples = ", ".join(
        f"{row['sample_id']}→{float(row['predicted_loss_rounded_1dp_W_per_m3']):.1f}"
        for row in summary["q4"]["specified_samples"]
    )
    rows = [
        ("E001", "Q1 分组 OOF Macro-F1=1.000；不等于附件二真实准确率", "q1/oof_metrics.json", "python src/verify_s5_outputs.py --config experiments/s2_frozen_config.json", "YES", "12,400 条分组 OOF"),
        ("E002", "Q1 相位/幅值一致率及留一材料最低 Macro-F1 均为 1.000", "q1/invariance_metrics.csv; q1/leave_one_material_out.csv", "同上", "YES", "支持保留 shape-only Logistic"),
        ("E003", f"附件二冻结分类共 80 条；正文指定样本编码：{q1_samples}", "q1/attachment2_predictions.csv", "python src/run_s5_predictions.py --config experiments/s2_frozen_config.json", "YES", "类别编码：1正弦、2三角、3梯形；无测试真值"),
        ("E004", "Q2 二次温度修正 OOF RMSLE 0.360678→0.202567，改善 43.84%", "q2/quadratic_temperature_metrics.json; q2/quadratic_temperature_parameters.csv", "python src/verify_s5_outputs.py --config experiments/s2_frozen_config.json", "YES", "材料1、正弦波；同折比较"),
        ("E005", "Q2 在 25°C/90°C 留一温度压力改善，但 50°C 局部劣化", "q2/leave_one_temperature_comparison.csv; q2/sensitivity_metrics.json", "同上", "YES", "须在局限性中披露"),
        ("E006", "Q3 两两交互模型 OOF log-RMSE 0.343416→0.326667，改善 4.88%", "q3/interaction_metrics.json", "同上", "YES", "仅解释为调整后关联"),
        ("E007", "Q3 完成 500/500 簇 Bootstrap；40 个两两对比中 27 个符号稳定度≥0.90", "q3/bootstrap_metrics.json; q3/bootstrap_pairwise_interaction_intervals.csv", "同上", "YES", "共同支持仅 14.24%，不得作因果表述"),
        ("E008", "Q4 HGB 同折优于 Ridge；最终采用 18 特征 amplitude_and_condition 变体", "q4/hgb_metrics.json; q4/ablation_metrics.json", "同上", "YES", "最终 OOF RMSLE=0.070941"),
        ("E009", "Q4 最终模型 12,400 条 OOF 预测可逐行追溯", "q4/final_oof_predictions.csv", "同上", "YES", "论文精度只来自附件一分组 OOF"),
        ("E010", "Q4 低 Bm OOF RMSLE=0.08328；LOMO/LOTO 最大 RMSLE 分别约 0.37982/0.57496", "q4/stress_metrics.json; q4/leave_one_level_out_metrics.csv", "同上", "YES", "属于外推压力，不是主 OOF"),
        ("E011", f"附件三冻结预测共 400 条；正文指定样本（W/m³）：{q4_samples}", "q4/attachment3_predictions.csv", "python src/run_s5_predictions.py --config experiments/s2_frozen_config.json", "YES", "保留 1 位小数；无测试真值"),
        ("E012", "Q5 主 OOF Pareto 点 118 个，折支持区域 27 个", "q5/robustness_metrics.json; q5/final_oof_pareto.csv", "同上", "YES", "只在观测联合可行域内解释"),
        ("E013", "Q5 完成 500/500 condition_group Bootstrap，稳定区域 42 个", "q5/bootstrap_region_stability.csv; q5/robustness_metrics.json", "同上", "YES", "固定交叉拟合候选表的重采样稳定性"),
        ("E014", "Q5 主模型/观测区域 Jaccard=0.39535<0.50，禁止唯一推荐", "q5/robustness_metrics.json", "同上", "YES", "敏感性结果不得覆盖主门槛失败"),
        ("E015", "附件四副本的第2列前80项与 Q1 一致，第3列400项与 Q4 一致", "submission/附件四（Excel表）.xlsx", "同上", "YES", "原始附件四哈希保持不变"),
    ]
    lines = [
        "# Result Registry",
        "",
        "> 状态：S5 已独立核验并冻结，等待 G5；论文和正式绘图只能引用本目录。",
        "",
        "| Evidence ID | Claim | Source File | Script / Command | Verified | Notes |",
        "|---|---|---|---|---|---|",
    ]
    lines.extend("| " + " | ".join(value.replace("|", "\\|") for value in row) + " |" for row in rows)
    lines.extend(
        [
            "",
            "## Global limitations",
            "",
            "- 附件二、附件三没有公开真值，正式预测不能作为新的精度证据。",
            "- Q3 结果是控制已建模协变量后的关联，不是因果效应。",
            "- Q4 的未见材料/温度外推压力明显高于主 OOF，不能宣称无条件域外泛化。",
            "- Q5 不满足唯一推荐门槛，只能报告 Pareto 权衡、稳定区域和端点。",
            "- G5 通过前，上述结果仍不得宣称为最终提交已批准。",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    args = parse_args()
    config_path = args.config.resolve()
    run = ExperimentRun(
        EXPERIMENT_ID,
        config_path,
        descriptor_path=DESCRIPTOR,
        evidence_root=RUN_ROOT,
        stage="S5",
    )
    checks: list[dict[str, Any]] = []
    try:
        config = load_frozen_config(config_path)
        run.record_contract_hashes(verify_contract_registry(config))
        required_inputs = (
            Path(__file__).resolve(),
            Path(__file__).with_name("run_s5_predictions.py").resolve(),
            TEST2_FILE,
            TEST3_FILE,
            TEMPLATE_FILE,
            Q1_MODEL_FILE,
            Q4_MODEL_FILE,
            Q1_RUN_MANIFEST,
            Q4_METRICS_FILE,
            RAW_Q1_PREDICTIONS,
            RAW_Q4_PREDICTIONS,
            RAW_SUBMISSION,
            RAW_FREEZE,
            RAW_SUMMARY,
            S4_VERIFICATION,
            RAW_S5_ROOT / "runs" / "EXP-S5-PRED-001" / "run_manifest.json",
        )
        for path in required_inputs:
            run.record_input(path)
        for source_relative in PROMOTION_MAP:
            run.record_input(PROJECT_ROOT / source_relative)

        input_hashes = verify_raw_inputs(config)
        record_check(checks, "official_input_hashes_match_frozen_config", len(input_hashes) == 5, input_hashes)
        authorization = assert_g4_pass()
        record_check(checks, "g4_pass_and_s5_authorized", authorization["reviewed_commit"] == G4_REVIEWED_COMMIT, authorization)

        s4_report = json.loads(S4_VERIFICATION.read_text(encoding="utf-8"))
        record_check(checks, "s4_independent_verification_23_of_23", s4_report.get("check_count") == 23 and s4_report.get("failed_count") == 0, s4_report)
        q1_manifest = json.loads(Q1_RUN_MANIFEST.read_text(encoding="utf-8"))
        expected_q1_hash = q1_manifest["output_hashes"]["results/raw/s3/q1/logistic_full_model.joblib"]
        record_check(checks, "q1_model_hash", sha256_file(Q1_MODEL_FILE) == expected_q1_hash, expected_q1_hash)
        q4_metrics = json.loads(Q4_METRICS_FILE.read_text(encoding="utf-8"))
        record_check(checks, "q4_model_hash", sha256_file(Q4_MODEL_FILE) == q4_metrics["final_full_model_sha256"], q4_metrics["final_full_model_sha256"])

        prediction_manifest = json.loads((RAW_S5_ROOT / "runs" / "EXP-S5-PRED-001" / "run_manifest.json").read_text(encoding="utf-8"))
        record_check(checks, "prediction_manifest_pass", prediction_manifest.get("status") == "PASS" and prediction_manifest.get("exit_code") == 0, prediction_manifest.get("status"))
        record_check(checks, "prediction_implementation_clean", prediction_manifest.get("implementation_git_dirty_at_finish") is False, prediction_manifest.get("implementation_git_dirty_at_finish"))
        record_check(checks, "prediction_stage_s5", prediction_manifest.get("stage") == "S5", prediction_manifest.get("stage"))

        plateau = float(config["feature_contract"]["plateau_relative_slope_threshold"])
        test2 = load_test_features(TEST2_FILE, "test2", plateau)
        test3 = load_test_features(TEST3_FILE, "test3", plateau)
        expected_q1 = predict_q1(test2, joblib.load(Q1_MODEL_FILE))
        expected_q4 = predict_q4(test3, joblib.load(Q4_MODEL_FILE))
        actual_q1 = pd.read_csv(RAW_Q1_PREDICTIONS).sort_values("sample_id").reset_index(drop=True)
        actual_q4 = pd.read_csv(RAW_Q4_PREDICTIONS).sort_values("sample_id").reset_index(drop=True)

        record_check(checks, "q1_row_and_id_coverage", len(actual_q1) == 80 and list(actual_q1["sample_id"]) == list(range(1, 81)), len(actual_q1))
        record_check(checks, "q1_labels_recomputed", list(actual_q1["predicted_label"]) == list(expected_q1["predicted_label"]), "80/80")
        record_check(checks, "q1_codes_recomputed", list(actual_q1["class_code"].astype(int)) == list(expected_q1["class_code"].astype(int)), "80/80")
        record_check(checks, "q1_label_code_consistency", all(CLASS_ENCODING[label] == int(code) for label, code in zip(actual_q1["predicted_label"], actual_q1["class_code"])), "80/80")
        for column in ("probability_sine", "probability_triangle", "probability_trapezoid"):
            compare_numeric(f"q1_{column}_recomputed", actual_q1[column], expected_q1[column], atol=1e-12, checks=checks)
        probability_sum = actual_q1[["probability_sine", "probability_triangle", "probability_trapezoid"]].sum(axis=1)
        record_check(checks, "q1_probability_sum", np.allclose(probability_sum, 1.0, atol=1e-12), float(np.max(np.abs(probability_sum - 1.0))))
        record_check(checks, "q1_waveform_lineage", list(actual_q1["waveform_sha256"]) == list(expected_q1["waveform_sha256"]), "80/80")

        record_check(checks, "q4_row_and_id_coverage", len(actual_q4) == 400 and list(actual_q4["sample_id"]) == list(range(1, 401)), len(actual_q4))
        compare_numeric("q4_log_prediction_recomputed", actual_q4["prediction_log"], expected_q4["prediction_log"], atol=1e-12, rtol=5e-12, checks=checks)
        compare_numeric("q4_loss_prediction_recomputed", actual_q4["predicted_loss_W_per_m3"], expected_q4["predicted_loss_W_per_m3"], atol=1e-9, rtol=5e-12, checks=checks)
        compare_numeric("q4_rounded_prediction_recomputed", actual_q4["predicted_loss_rounded_1dp_W_per_m3"], expected_q4["predicted_loss_rounded_1dp_W_per_m3"], atol=0.0, checks=checks)
        record_check(checks, "q4_predictions_positive_finite", np.isfinite(actual_q4["predicted_loss_W_per_m3"]).all() and (actual_q4["predicted_loss_W_per_m3"] > 0).all(), [float(actual_q4["predicted_loss_W_per_m3"].min()), float(actual_q4["predicted_loss_W_per_m3"].max())])
        record_check(checks, "q4_waveform_lineage", list(actual_q4["waveform_sha256"]) == list(expected_q4["waveform_sha256"]), "400/400")
        verify_submission(actual_q1, actual_q4, checks)
        record_check(checks, "original_template_unchanged", sha256_file(TEMPLATE_FILE) == input_hashes["src/附件四（Excel表）.xlsx"], sha256_file(TEMPLATE_FILE))

        summary = json.loads(RAW_SUMMARY.read_text(encoding="utf-8"))
        freeze = json.loads(RAW_FREEZE.read_text(encoding="utf-8"))
        record_check(checks, "summary_policy_no_test_tuning", summary.get("selection_or_tuning_with_test_attachments") is False, summary.get("prediction_policy"))
        record_check(checks, "freeze_model_hashes", freeze["q1"]["model_sha256"] == sha256_file(Q1_MODEL_FILE) and freeze["q4"]["model_sha256"] == sha256_file(Q4_MODEL_FILE), {"q1": freeze["q1"]["model_sha256"], "q4": freeze["q4"]["model_sha256"]})
        record_check(checks, "freeze_status_pending_g5", freeze.get("status") == "FROZEN_PENDING_G5", freeze.get("status"))
        q5_metrics = json.loads((PROJECT_ROOT / "results" / "raw" / "main" / "q5" / "robustness_metrics.json").read_text(encoding="utf-8"))
        record_check(checks, "q5_unique_recommendation_still_forbidden", q5_metrics.get("unique_recommendation_authorized") is False and q5_metrics.get("observed_region_jaccard_pass") is False, {"jaccard": q5_metrics.get("observed_region_jaccard")})

        provenance = promote_whitelist(checks)
        registry_path = VERIFIED_ROOT / "result_registry.md"
        registry_path.write_text(build_registry(summary), encoding="utf-8", newline="\n")
        registry_lines = registry_path.read_text(encoding="utf-8").splitlines()
        evidence_ids = [
            line.split("|")[1].strip()
            for line in registry_lines
            if line.startswith("| E") and line.split("|")[1].strip()[1:].isdigit()
        ]
        record_check(
            checks,
            "result_registry_has_15_evidence_ids",
            evidence_ids == [f"E{index:03d}" for index in range(1, 16)],
            evidence_ids,
        )

        report = {
            "status": "PASS",
            "check_count": len(checks),
            "failed_count": sum(item["status"] != "PASS" for item in checks),
            "checks": checks,
            "g4_authorization": authorization,
            "prediction_manifest_git_commit": prediction_manifest["git_commit"],
            "test_recompute_policy": "in-memory verification only; raw official prediction artifacts were not overwritten",
            "verified_file_count_excluding_registry_and_reports": len(provenance),
        }
        report_path = VERIFIED_ROOT / "verification_report.json"
        write_json(report_path, report)
        provenance_path = VERIFIED_ROOT / "provenance.json"
        write_json(
            provenance_path,
            {
                "status": "PASS",
                "promotion_policy": "explicit_whitelist_after_all_pre_promotion_checks_pass",
                "files": provenance,
                "registry_sha256": sha256_file(registry_path),
                "verification_report_sha256": sha256_file(report_path),
            },
        )
        for path in [*(VERIFIED_ROOT / relative for relative in PROMOTION_MAP.values()), registry_path, report_path, provenance_path]:
            run.record_output(path)
        run.log(json.dumps({"check_count": len(checks), "verified_files": len(provenance)}, sort_keys=True))
        run.finish("PASS")
    except BaseException as exc:
        run.fail(exc)
        raise


if __name__ == "__main__":
    main()
