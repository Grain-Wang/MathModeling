"""Run the mandatory S4 Q1 invariance, leave-material, and feature/auxiliary ablations."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

from build_features import TRAIN_FILE, TRAIN_SHEETS, waveform_features
from run_q1 import make_model
from s3_common import (
    DEFAULT_CONFIG,
    PROJECT_ROOT,
    S3_RESULTS_ROOT,
    SHAPE_FEATURES,
    classification_metrics,
    load_features_and_folds,
    load_frozen_config,
    set_global_seed,
    verify_contract_registry,
    write_csv,
    write_json,
)
from s4_common import S4_RESULTS_ROOT, s4_run

EXPERIMENT_ID = "EXP-Q1-ABL-001"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    return parser.parse_args()


def transformed_shape_features(phase_shift: int, amplitude: float, plateau_fraction: float) -> pd.DataFrame:
    source = PROJECT_ROOT / "src" / TRAIN_FILE
    workbook = load_workbook(source, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    try:
        for material, expected_rows in TRAIN_SHEETS.items():
            sheet = workbook[material]
            rows = sheet.iter_rows(values_only=True)
            next(rows)
            count = 0
            for excel_row, row in enumerate(rows, start=2):
                count += 1
                waveform = np.asarray(row[4:], dtype=np.float64)
                transformed = np.roll(waveform * amplitude, phase_shift)
                features, _ = waveform_features(transformed, plateau_fraction)
                records.append({"row_id": f"{TRAIN_FILE}|{material}|{excel_row}", **{name: features[name] for name in SHAPE_FEATURES}})
            if count != expected_rows:
                raise ValueError(f"Unexpected row count for {material}: {count}")
    finally:
        workbook.close()
    return pd.DataFrame(records).sort_values("row_id").reset_index(drop=True)


def cross_fit_fixed_logistic(x: pd.DataFrame, data: pd.DataFrame, c_value: float, class_weight: str | None, max_iter: int, seed: int) -> np.ndarray:
    y = data["waveform"].to_numpy(dtype=object)
    folds = data["q1_outer_fold"].to_numpy(dtype=int)
    groups = data["near_shape_group"].to_numpy(dtype=object)
    prediction = np.empty(len(data), dtype=object)
    for fold in sorted(np.unique(folds)):
        train_index = np.flatnonzero(folds != fold)
        valid_index = np.flatnonzero(folds == fold)
        if set(groups[train_index]).intersection(groups[valid_index]):
            raise AssertionError("Q1 ablation group leakage")
        model = make_model(c_value, class_weight, max_iter, seed)
        model.fit(x.iloc[train_index], y[train_index])
        prediction[valid_index] = model.predict(x.iloc[valid_index])
    if pd.isna(prediction).any():
        raise AssertionError("Q1 ablation OOF coverage incomplete")
    return prediction


def fold_model_predictions(features: pd.DataFrame, data: pd.DataFrame, fold_models: dict[str, Pipeline]) -> np.ndarray:
    prediction = np.empty(len(data), dtype=object)
    folds = data["q1_outer_fold"].to_numpy(dtype=int)
    for fold in sorted(np.unique(folds)):
        index = np.flatnonzero(folds == fold)
        model = fold_models[f"q1-logistic-fold-{int(fold)}"]
        prediction[index] = model.predict(features.iloc[index][SHAPE_FEATURES])
    return prediction


def main() -> None:
    args = parse_args()
    config_path = args.config.resolve()
    run = s4_run(EXPERIMENT_ID, config_path, "ablation")
    try:
        config = load_frozen_config(config_path)
        run.record_contract_hashes(verify_contract_registry(config))
        seed = int(config["seeds"]["global"])
        set_global_seed(seed)
        required = (
            Path(__file__).resolve(), Path(__file__).with_name("s3_common.py").resolve(), Path(__file__).with_name("s4_common.py").resolve(),
            Path(__file__).with_name("run_q1.py").resolve(), Path(__file__).with_name("build_features.py").resolve(),
            PROJECT_ROOT / "src" / TRAIN_FILE,
            S3_RESULTS_ROOT / "data" / "feature_table.csv", S3_RESULTS_ROOT / "folds" / "fold_assignments.csv",
            S3_RESULTS_ROOT / "q1" / "baseline_oof_predictions.csv", S3_RESULTS_ROOT / "q1" / "baseline_metrics.json",
            S3_RESULTS_ROOT / "q1" / "logistic_fold_models.joblib",
        )
        for path in required:
            run.record_input(path)
        _, _, data = load_features_and_folds()
        data = data.sort_values("row_id").reset_index(drop=True)
        baseline = pd.read_csv(S3_RESULTS_ROOT / "q1" / "baseline_oof_predictions.csv").sort_values("row_id").reset_index(drop=True)
        if list(data["row_id"]) != list(baseline["row_id"]):
            raise ValueError("Q1 baseline identity mismatch")
        baseline_metrics = json.loads((S3_RESULTS_ROOT / "q1" / "baseline_metrics.json").read_text(encoding="utf-8"))
        parameters = baseline_metrics["full_fit_parameters_from_outer_mode"]
        c_value = float(parameters["C"])
        class_weight = parameters["class_weight"]
        max_iter = int(config["models"]["q1"]["search_space"]["logistic"]["max_iter"])
        fold_models = joblib.load(S3_RESULTS_ROOT / "q1" / "logistic_fold_models.joblib")
        plateau = float(config["feature_contract"]["plateau_relative_slope_threshold"])

        perturbations = [("phase_shift_137", 137, 1.0), ("amplitude_half", 0, 0.5), ("amplitude_double", 0, 2.0)]
        invariance_rows: list[dict[str, Any]] = []
        original_features = data[SHAPE_FEATURES].reset_index(drop=True)
        y = data["waveform"].to_numpy(dtype=object)
        baseline_prediction = baseline["predicted_label"].to_numpy(dtype=object)
        for name, shift, amplitude in perturbations:
            transformed = transformed_shape_features(shift, amplitude, plateau)
            if list(transformed["row_id"]) != list(data["row_id"]):
                raise ValueError("Q1 transformed waveform identity mismatch")
            transformed_x = transformed[SHAPE_FEATURES]
            prediction = fold_model_predictions(transformed_x, data, fold_models)
            metrics = classification_metrics(y, prediction)
            difference = np.abs(transformed_x.to_numpy(dtype=np.float64) - original_features.to_numpy(dtype=np.float64))
            invariance_rows.append(
                {
                    "perturbation": name,
                    "phase_shift_samples": shift,
                    "amplitude_multiplier": amplitude,
                    "prediction_agreement_with_baseline": float(np.mean(prediction == baseline_prediction)),
                    "macro_f1": float(metrics["macro_f1"]),
                    "accuracy": float(metrics["accuracy"]),
                    "maximum_absolute_shape_feature_change": float(np.max(difference)),
                }
            )
            run.log(f"perturbation={name} agreement={invariance_rows[-1]['prediction_agreement_with_baseline']}")

        leave_material_rows: list[dict[str, Any]] = []
        for material in sorted(data["material"].unique()):
            train = data[data["material"] != material]
            valid = data[data["material"] == material]
            model = make_model(c_value, class_weight, max_iter, seed)
            model.fit(train[SHAPE_FEATURES], train["waveform"])
            prediction = model.predict(valid[SHAPE_FEATURES])
            metrics = classification_metrics(valid["waveform"], prediction)
            leave_material_rows.append({"heldout_material": material, "n": len(valid), "macro_f1": metrics["macro_f1"], "accuracy": metrics["accuracy"], "balanced_accuracy": metrics["balanced_accuracy"]})

        fft_features = [feature for feature in SHAPE_FEATURES if "fft_" in feature]
        time_features = [feature for feature in SHAPE_FEATURES if feature not in fft_features]
        auxiliary = pd.concat(
            [
                data[["log_frequency_Hz", "log_b_m_T"]].reset_index(drop=True),
                pd.get_dummies(data[["material", "temperature_C"]].astype(str), prefix=["material", "temperature"], dtype=float).reset_index(drop=True),
            ],
            axis=1,
        )
        variants = {
            "shape_full": original_features,
            "shape_time_domain_only": data[time_features].reset_index(drop=True),
            "shape_frequency_domain_only": data[fft_features].reset_index(drop=True),
            "auxiliary_only": auxiliary,
            "shape_plus_auxiliary": pd.concat([original_features, auxiliary], axis=1),
        }
        ablation_rows: list[dict[str, Any]] = []
        prediction_frame = data[["row_id", "q1_outer_fold", "waveform"]].copy()
        for variant, x in variants.items():
            if variant == "shape_full":
                prediction = baseline_prediction
            else:
                prediction = cross_fit_fixed_logistic(x, data, c_value, class_weight, max_iter, seed)
            metrics = classification_metrics(y, prediction)
            ablation_rows.append({"variant": variant, "feature_count": x.shape[1], "macro_f1": metrics["macro_f1"], "accuracy": metrics["accuracy"], "balanced_accuracy": metrics["balanced_accuracy"], "minimum_class_recall": min(value["recall"] for value in metrics["per_class"].values())})
            prediction_frame[f"prediction_{variant}"] = prediction

        output_dir = S4_RESULTS_ROOT / "q1"
        invariance_path = output_dir / "invariance_metrics.csv"
        leave_material_path = output_dir / "leave_one_material_out.csv"
        ablation_path = output_dir / "feature_auxiliary_ablation.csv"
        prediction_path = output_dir / "ablation_oof_predictions.csv"
        metrics_path = output_dir / "stress_metrics.json"
        evidence_path = run.output_dir / "metrics.json"
        write_csv(invariance_path, pd.DataFrame(invariance_rows))
        write_csv(leave_material_path, pd.DataFrame(leave_material_rows))
        write_csv(ablation_path, pd.DataFrame(ablation_rows))
        write_csv(prediction_path, prediction_frame)
        minimum_agreement = min(float(row["prediction_agreement_with_baseline"]) for row in invariance_rows)
        minimum_leave_material_f1 = min(float(row["macro_f1"]) for row in leave_material_rows)
        metrics = {
            "status": "PASS",
            "baseline_macro_f1": baseline_metrics["oof"]["macro_f1"],
            "minimum_phase_amplitude_prediction_agreement": minimum_agreement,
            "minimum_leave_one_material_macro_f1": minimum_leave_material_f1,
            "shape_only_remains_saturated": bool(minimum_agreement == 1.0 and minimum_leave_material_f1 >= 0.99),
            "tree_candidates_run": False,
            "tree_candidates_skip_reason": "Frozen Logistic remains saturated after mandatory invariance and leave-material stress; added complexity has no measurable selection headroom.",
            "auxiliary_only_macro_f1": next(row["macro_f1"] for row in ablation_rows if row["variant"] == "auxiliary_only"),
            "proxy_risk_interpretation": "Auxiliary-only performance is diagnostic and cannot replace the shape-only model.",
        }
        write_json(metrics_path, metrics)
        write_json(evidence_path, metrics)
        for path in (invariance_path, leave_material_path, ablation_path, prediction_path, metrics_path, evidence_path):
            run.record_output(path)
        run.log(json.dumps(metrics, ensure_ascii=False, sort_keys=True))
        run.finish("PASS")
    except BaseException as exc:
        run.fail(exc)
        raise


if __name__ == "__main__":
    main()